"""Khảo sát — biểu mẫu kiểu Google Forms gửi tới các nhóm user, có hạn chót và thống kê.

Hai lớp quyền khác nhau, đừng trộn:
- Tạo / sửa / xem kết quả: mã quyền (`surveys.create`, `surveys.view_all`) gán
  qua màn Phân quyền theo nhóm.
- TRẢ LỜI: chỉ cần có tên trong `survey_recipients`. Đó là dữ liệu của từng
  khảo sát (giống người được giao duyệt đơn), không phải quyền — gate bằng mã
  quyền thì gửi khảo sát cho nhóm chưa được tick `menu.surveys` là cả nhóm
  không trả lời được mà không ai hay (xem docs/DESIGN.md mục Phân quyền).
"""
import io
import json
import sqlite3
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.api.quiz import _download_headers
from backend.core.deps import get_current_staff, require_feature
from backend.core.enums import StaffRole
from backend.database import get_db, write_audit
from backend.schemas.surveys import ResponseIn, SurveyIn
from backend.services import survey_service as svc

router = APIRouter(prefix="/api/surveys", tags=["surveys"])


# ── Tiện ích ──────────────────────────────────────────────────────────────────
def _co_quyen(db: sqlite3.Connection, current: dict, code: str) -> bool:
    """Admin luôn có; còn lại tra quyền gán qua nhóm (cùng SQL với require_feature)."""
    if current.get("role") == StaffRole.ADMIN:
        return True
    return bool(db.execute(
        """SELECT 1 FROM group_features gf
           JOIN group_members gm ON gm.group_id = gf.group_id
           JOIN user_groups g ON g.id = gm.group_id AND g.is_active = 1
           WHERE gm.staff_id = ? AND gf.feature_code = ? LIMIT 1""",
        (current["id"], code),
    ).fetchone())


def _lay(db: sqlite3.Connection, survey_id: int) -> dict:
    row = db.execute(
        """SELECT s.*, u.full_name AS created_by_name FROM surveys s
             LEFT JOIN user_tttt u ON u.id = s.created_by WHERE s.id = ?""",
        (survey_id,),
    ).fetchone()
    if not row:
        raise HTTPException(404, "Không tìm thấy khảo sát")
    return dict(row)


def _la_chu(current: dict, s: dict) -> bool:
    # Người tạo là dữ liệu của hồ sơ; vai admin là siêu quyền cố ý (DESIGN.md).
    return s["created_by"] == current["id"] or current["role"] == StaffRole.ADMIN


def _chan_neu_khong_phai_chu(current: dict, s: dict) -> None:
    if not _la_chu(current, s):
        raise HTTPException(403, "Chỉ người tạo khảo sát mới sửa được")


def _duoc_xem_ket_qua(db, current: dict, s: dict) -> bool:
    return _la_chu(current, s) or _co_quyen(db, current, "surveys.view_all")


def _so_tra_loi(db, survey_id: int) -> int:
    return db.execute(
        "SELECT COUNT(*) c FROM survey_responses WHERE survey_id = ?", (survey_id,)
    ).fetchone()["c"]


def _ghi_cau_hoi(db, survey_id: int, questions) -> None:
    db.execute("DELETE FROM survey_questions WHERE survey_id = ?", (survey_id,))
    db.executemany(
        """INSERT INTO survey_questions (survey_id, order_no, qtype, title, description,
                   required, options, scale_min, scale_max, scale_min_label, scale_max_label)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        [(survey_id, i, q.qtype, q.title, q.description, int(q.required),
          json.dumps(q.options, ensure_ascii=False), q.scale_min, q.scale_max,
          q.scale_min_label, q.scale_max_label)
         for i, q in enumerate(questions, start=1)],
    )


def _ghi_nhom(db, survey_id: int, group_ids: list[int]) -> None:
    if group_ids:
        ok = {r["id"] for r in db.execute(
            f"SELECT id FROM user_groups WHERE id IN ({','.join('?' * len(group_ids))})",
            group_ids).fetchall()}
        thieu = sorted(set(group_ids) - ok)
        if thieu:
            raise HTTPException(400, f"Không tìm thấy nhóm user id {thieu}")
    db.execute("DELETE FROM survey_target_groups WHERE survey_id = ?", (survey_id,))
    db.executemany(
        "INSERT INTO survey_target_groups (survey_id, group_id) VALUES (?,?)",
        [(survey_id, g) for g in group_ids],
    )


def _dieu_kien_phat_hanh(db, s: dict) -> None:
    """Chặn phát hành khảo sát không ai trả lời được / không có gì để trả lời."""
    if not db.execute("SELECT 1 FROM survey_questions WHERE survey_id = ?", (s["id"],)).fetchone():
        raise HTTPException(400, "Khảo sát chưa có câu hỏi nào")
    if not db.execute("SELECT 1 FROM survey_target_groups WHERE survey_id = ?", (s["id"],)).fetchone():
        raise HTTPException(400, "Chưa chọn nhóm user nhận khảo sát")
    if not s["deadline"]:
        raise HTTPException(400, "Chưa đặt hạn chót cho khảo sát")
    if s["deadline"] <= svc.now_str():
        raise HTTPException(400, "Hạn chót phải sau thời điểm hiện tại")


def _dem_nguoi_nhan(db, survey_id: int) -> int:
    return db.execute(
        "SELECT COUNT(*) c FROM survey_recipients WHERE survey_id = ?", (survey_id,)
    ).fetchone()["c"]


# ── Danh sách ─────────────────────────────────────────────────────────────────
# Các route tĩnh (/mine, /manage, /groups) PHẢI đứng trên /{survey_id}: FastAPI
# khớp theo thứ tự đăng ký, để sau thì "mine" rơi vào {survey_id} và ăn 422.
@router.get("/mine")
def my_surveys(
    current: dict = Depends(get_current_staff),
    db: sqlite3.Connection = Depends(get_db),
):
    """Khảo sát gửi tới tôi (đã phát hành). Không đòi mã quyền — xem docstring module."""
    rows = db.execute(
        """SELECT s.id, s.title, s.description, s.status, s.start_at, s.deadline,
                  s.allow_edit, s.is_anonymous, u.full_name AS created_by_name,
                  x.submitted_at, x.updated_at,
                  (SELECT COUNT(*) FROM survey_questions q WHERE q.survey_id = s.id) AS question_count
             FROM survey_recipients r
             JOIN surveys s ON s.id = r.survey_id AND s.status <> 'draft'
             LEFT JOIN user_tttt u ON u.id = s.created_by
             LEFT JOIN survey_responses x ON x.survey_id = s.id AND x.staff_id = r.staff_id
            WHERE r.staff_id = ?
            ORDER BY s.deadline DESC""",
        (current["id"],),
    ).fetchall()
    now = svc.now_str()
    out = []
    for r in rows:
        d = dict(r)
        d["state"] = svc.trang_thai(d, now)
        d["responded"] = d["submitted_at"] is not None
        d["allow_edit"] = bool(d["allow_edit"])
        d["is_anonymous"] = bool(d["is_anonymous"])
        out.append(d)
    # Việc cần làm ngay lên đầu: đang mở + chưa trả lời, hạn gần nhất trước.
    out.sort(key=lambda d: (not (d["state"] == "open" and not d["responded"]),
                            d["deadline"] or "9999"))
    return out


@router.get("/manage")
def manage_surveys(
    scope: str = Query("mine", pattern="^(mine|all)$"),
    current: dict = Depends(require_feature("menu.surveys")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Khảo sát tôi tạo; `scope=all` (cần surveys.view_all) là của mọi người."""
    if scope == "all" and not _co_quyen(db, current, "surveys.view_all"):
        raise HTTPException(403, "Không có quyền xem khảo sát của người khác")
    sql = """SELECT s.id, s.title, s.status, s.start_at, s.deadline, s.created_at,
                    s.published_at, s.is_anonymous, s.created_by, u.full_name AS created_by_name,
                    (SELECT COUNT(*) FROM survey_questions q WHERE q.survey_id = s.id) AS question_count,
                    (SELECT COUNT(*) FROM survey_recipients r WHERE r.survey_id = s.id) AS recipient_count,
                    (SELECT COUNT(*) FROM survey_responses x WHERE x.survey_id = s.id) AS response_count
               FROM surveys s LEFT JOIN user_tttt u ON u.id = s.created_by"""
    params: list = []
    if scope == "mine":
        sql += " WHERE s.created_by = ?"
        params.append(current["id"])
    elif current["role"] != StaffRole.ADMIN:
        # Bản nháp của người khác không thuộc "xem kết quả" (cùng lý do với get_survey)
        sql += " WHERE (s.status <> 'draft' OR s.created_by = ?)"
        params.append(current["id"])
    sql += " ORDER BY s.created_at DESC"
    now = svc.now_str()
    out = []
    for r in db.execute(sql, params).fetchall():
        d = dict(r)
        d["state"] = svc.trang_thai(d, now)
        d["is_owner"] = _la_chu(current, d)
        out.append(d)
    return out


@router.get("/groups")
def target_groups(
    _: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Nhóm user chọn làm người nhận. Đếm đúng cách dong_bo_nguoi_nhan() sẽ chốt."""
    rows = db.execute(
        """SELECT g.id, g.name, g.description,
                  COUNT(DISTINCT CASE WHEN u.is_active = 1 AND IFNULL(u.is_deleted, 0) = 0
                                      THEN u.id END) AS member_count
             FROM user_groups g
             LEFT JOIN group_members gm ON gm.group_id = g.id
             LEFT JOIN user_tttt u ON u.id = gm.staff_id
            WHERE g.is_active = 1
            GROUP BY g.id ORDER BY g.name"""
    ).fetchall()
    return [dict(r) for r in rows]


# ── Tạo / sửa ─────────────────────────────────────────────────────────────────
@router.post("", status_code=201)
def create_survey(
    body: SurveyIn,
    current: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    now = svc.now_str()
    sid = db.execute(
        """INSERT INTO surveys (title, description, status, is_anonymous, allow_edit,
                               start_at, deadline, created_by, created_at, updated_at)
           VALUES (?,?, 'draft', ?,?,?,?,?,?,?)""",
        (body.title, body.description, int(body.is_anonymous), int(body.allow_edit),
         body.start_at, body.deadline, current["id"], now, now),
    ).lastrowid
    _ghi_cau_hoi(db, sid, body.questions)
    _ghi_nhom(db, sid, body.group_ids)
    write_audit(db, current["id"], "survey.create", "survey", sid, body.title)
    db.commit()
    return {"id": sid}


@router.get("/{survey_id}")
def get_survey(
    survey_id: int,
    current: dict = Depends(require_feature("menu.surveys")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Định nghĩa đầy đủ cho màn soạn thảo."""
    s = _lay(db, survey_id)
    if not _duoc_xem_ket_qua(db, current, s):
        raise HTTPException(403, "Không có quyền xem khảo sát này")
    # surveys.view_all là xem KẾT QUẢ — bản nháp của người khác chưa có gì để xem
    if s["status"] == "draft" and not _la_chu(current, s):
        raise HTTPException(404, "Không tìm thấy khảo sát")
    s["is_anonymous"] = bool(s["is_anonymous"])
    s["allow_edit"] = bool(s["allow_edit"])
    s["state"] = svc.trang_thai(s)
    s["group_ids"] = [r["group_id"] for r in db.execute(
        "SELECT group_id FROM survey_target_groups WHERE survey_id = ?", (survey_id,)).fetchall()]
    s["questions"] = svc.doc_cau_hoi(db, survey_id)
    s["response_count"] = _so_tra_loi(db, survey_id)
    s["recipient_count"] = _dem_nguoi_nhan(db, survey_id)
    # Có câu trả lời thì khoá phần câu hỏi — xem dau_van_tay_cau_hoi().
    s["locked"] = s["response_count"] > 0
    s["can_edit"] = _la_chu(current, s) and _co_quyen(db, current, "surveys.create")
    return s


@router.put("/{survey_id}")
def update_survey(
    survey_id: int,
    body: SurveyIn,
    current: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    s = _lay(db, survey_id)
    _chan_neu_khong_phai_chu(current, s)

    # ── Khoá câu hỏi khi đã có người trả lời ──
    locked = _so_tra_loi(db, survey_id) > 0
    if locked:
        cu = svc.dau_van_tay_cau_hoi(svc.doc_cau_hoi(db, survey_id))
        moi = svc.dau_van_tay_cau_hoi([q.model_dump() for q in body.questions])
        if cu != moi:
            raise HTTPException(
                409, "Đã có người trả lời nên không sửa được câu hỏi. "
                     "Chỉ đổi được tên, mô tả, thời hạn và nhóm nhận.")

    # ── Chế độ ẩn danh — khoá ở màn soạn thảo chưa đủ, gọi thẳng API vẫn đổi được ──
    # Tắt sau phát hành: người nhận đã thấy dòng "Ẩn danh" trên biểu mẫu, câu trả lời
    # họ nộp trong niềm tin đó sẽ hiện nguyên tên.
    # Bật khi đã có trả lời: chủ đã xem các bài có tên, người nộp SAU thấy "Ẩn danh"
    # rồi nộp — trừ các bài đã biết là ra đúng bài của họ.
    if s["is_anonymous"] and not body.is_anonymous and s["status"] != "draft":
        raise HTTPException(409, "Khảo sát đã phát hành ở chế độ ẩn danh — không tắt ẩn danh được")
    if locked and bool(s["is_anonymous"]) != body.is_anonymous:
        raise HTTPException(409, "Đã có người trả lời nên không đổi được chế độ ẩn danh")

    # Khảo sát đang chạy mà dời hạn chót về quá khứ = đóng lén, người chưa trả
    # lời mất lượt không một lời báo. Muốn dừng thì bấm "Đóng khảo sát".
    if (s["status"] == "published" and body.deadline != s["deadline"]
            and (not body.deadline or body.deadline <= svc.now_str())):
        raise HTTPException(400, "Hạn chót mới phải sau thời điểm hiện tại")
    if s["status"] == "published" and not body.group_ids:
        raise HTTPException(400, "Khảo sát đang phát hành phải có ít nhất một nhóm nhận")

    db.execute(
        """UPDATE surveys SET title = ?, description = ?, is_anonymous = ?, allow_edit = ?,
                              start_at = ?, deadline = ?, updated_at = ? WHERE id = ?""",
        (body.title, body.description, int(body.is_anonymous), int(body.allow_edit),
         body.start_at, body.deadline, svc.now_str(), survey_id),
    )
    if not locked:
        _ghi_cau_hoi(db, survey_id, body.questions)
        # Đếm lại SAU khi đã giữ khoá ghi: lần đếm ở trên chạy ngoài giao dịch, một bài
        # nộp chen vào giữa hai lúc sẽ mất sạch câu trả lời vì xoá câu hỏi kéo theo
        # survey_answers (FK CASCADE). Rollback thì bài đó còn nguyên.
        if _so_tra_loi(db, survey_id) > 0:
            db.rollback()
            raise HTTPException(409, "Vừa có người nộp câu trả lời — tải lại trang rồi sửa lại")
    else:
        # Khoá chỉ áp cho phần ảnh hưởng thống kê; mô tả và nhãn hai đầu thang
        # đo vẫn sửa được. Không ghi thì ô trên màn hình sửa được mà bấm Lưu
        # xong mất trắng. Vân tay đã khớp nên thứ tự câu trùng với DB.
        db.executemany(
            """UPDATE survey_questions SET description = ?, scale_min_label = ?, scale_max_label = ?
                WHERE survey_id = ? AND order_no = ?""",
            [(q.description, q.scale_min_label, q.scale_max_label, survey_id, i)
             for i, q in enumerate(body.questions, start=1)],
        )
    _ghi_nhom(db, survey_id, body.group_ids)

    added = removed = 0
    if s["status"] == "published":
        added, removed = svc.dong_bo_nguoi_nhan(db, survey_id)
        if _dem_nguoi_nhan(db, survey_id) == 0:
            raise HTTPException(400, "Các nhóm đã chọn không còn thành viên nào đang hoạt động")
    write_audit(db, current["id"], "survey.update", "survey", survey_id,
                f"{body.title} — người nhận +{added}/-{removed}")
    db.commit()
    return {"ok": True, "added": added, "removed": removed}


@router.post("/{survey_id}/publish")
def publish_survey(
    survey_id: int,
    current: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Phát hành bản nháp, hoặc mở lại khảo sát đã đóng (cần hạn chót mới)."""
    s = _lay(db, survey_id)
    _chan_neu_khong_phai_chu(current, s)
    if s["status"] == "published":
        raise HTTPException(409, "Khảo sát đã phát hành rồi")
    _dieu_kien_phat_hanh(db, s)

    added, _ = svc.dong_bo_nguoi_nhan(db, survey_id)
    n = _dem_nguoi_nhan(db, survey_id)
    if n == 0:
        db.rollback()
        raise HTTPException(400, "Các nhóm đã chọn không có thành viên nào đang hoạt động")
    now = svc.now_str()
    db.execute(
        """UPDATE surveys SET status = 'published', published_at = IFNULL(published_at, ?),
                              closed_at = NULL, updated_at = ? WHERE id = ?""",
        (now, now, survey_id),
    )
    write_audit(db, current["id"], "survey.publish", "survey", survey_id,
                f"{s['title']} — {n} người nhận, hạn {s['deadline']}")
    db.commit()
    return {"ok": True, "recipient_count": n}


@router.post("/{survey_id}/close")
def close_survey(
    survey_id: int,
    current: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    s = _lay(db, survey_id)
    _chan_neu_khong_phai_chu(current, s)
    if s["status"] != "published":
        raise HTTPException(409, "Chỉ đóng được khảo sát đang phát hành")
    now = svc.now_str()
    db.execute("UPDATE surveys SET status = 'closed', closed_at = ?, updated_at = ? WHERE id = ?",
               (now, now, survey_id))
    write_audit(db, current["id"], "survey.close", "survey", survey_id, s["title"])
    db.commit()
    return {"ok": True}


@router.delete("/{survey_id}")
def delete_survey(
    survey_id: int,
    current: dict = Depends(require_feature("surveys.create")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Xoá HẲN — câu hỏi, người nhận và mọi câu trả lời đi theo (FK CASCADE)."""
    s = _lay(db, survey_id)
    _chan_neu_khong_phai_chu(current, s)
    n = _so_tra_loi(db, survey_id)
    db.execute("DELETE FROM surveys WHERE id = ?", (survey_id,))
    write_audit(db, current["id"], "survey.delete", "survey", survey_id,
                f"{s['title']} — {n} câu trả lời")
    db.commit()
    return {"ok": True, "deleted_responses": n}


# ── Trả lời ───────────────────────────────────────────────────────────────────
@router.get("/{survey_id}/form")
def get_form(
    survey_id: int,
    current: dict = Depends(get_current_staff),
    db: sqlite3.Connection = Depends(get_db),
):
    """Biểu mẫu để trả lời. Người tạo / người xem kết quả mở được ở chế độ xem trước."""
    s = _lay(db, survey_id)
    is_rec = db.execute(
        "SELECT 1 FROM survey_recipients WHERE survey_id = ? AND staff_id = ?",
        (survey_id, current["id"]),
    ).fetchone() is not None
    if not is_rec and not _duoc_xem_ket_qua(db, current, s):
        raise HTTPException(403, "Khảo sát này không gửi tới bạn")
    if s["status"] == "draft" and not _la_chu(current, s):
        raise HTTPException(404, "Không tìm thấy khảo sát")

    mine = db.execute(
        "SELECT id, submitted_at, updated_at FROM survey_responses WHERE survey_id = ? AND staff_id = ?",
        (survey_id, current["id"]),
    ).fetchone()
    answers = {}
    if mine:
        answers = {str(r["question_id"]): json.loads(r["value"]) for r in db.execute(
            "SELECT question_id, value FROM survey_answers WHERE response_id = ?", (mine["id"],))}
    state = svc.trang_thai(s)
    can_submit = is_rec and state == "open" and (not mine or bool(s["allow_edit"]))
    return {
        "id": s["id"], "title": s["title"], "description": s["description"],
        "deadline": s["deadline"], "start_at": s["start_at"], "state": state,
        "is_anonymous": bool(s["is_anonymous"]), "allow_edit": bool(s["allow_edit"]),
        "created_by_name": s["created_by_name"],
        "questions": svc.doc_cau_hoi(db, survey_id),
        "is_recipient": is_rec, "responded": mine is not None,
        "submitted_at": mine["submitted_at"] if mine else None,
        "answers": answers, "can_submit": can_submit,
    }


@router.post("/{survey_id}/responses")
def submit_response(
    survey_id: int,
    body: ResponseIn,
    current: dict = Depends(get_current_staff),
    db: sqlite3.Connection = Depends(get_db),
):
    s = _lay(db, survey_id)
    if not db.execute(
        "SELECT 1 FROM survey_recipients WHERE survey_id = ? AND staff_id = ?",
        (survey_id, current["id"]),
    ).fetchone():
        raise HTTPException(403, "Khảo sát này không gửi tới bạn")
    state = svc.trang_thai(s)
    if state != "open":
        msg = {"scheduled": "Khảo sát chưa tới giờ mở", "expired": "Khảo sát đã quá hạn",
               "closed": "Khảo sát đã đóng"}.get(state, "Khảo sát không nhận câu trả lời")
        raise HTTPException(409, msg)

    mine = db.execute(
        "SELECT id FROM survey_responses WHERE survey_id = ? AND staff_id = ?",
        (survey_id, current["id"]),
    ).fetchone()
    if mine and not s["allow_edit"]:
        raise HTTPException(409, "Bạn đã trả lời khảo sát này rồi")

    questions = svc.doc_cau_hoi(db, survey_id)
    # Sửa khảo sát (khi chưa ai trả lời) là ghi lại câu hỏi với id MỚI. Người đang mở
    # biểu mẫu cũ bấm Gửi sẽ mang id cũ; kiem_tra_tra_loi bỏ qua id lạ → lưu một bài
    # rỗng, báo thành công, và từ đó câu hỏi bị khoá vì "đã có người trả lời".
    if {a.question_id for a in body.answers} - {q["id"] for q in questions}:
        raise HTTPException(409, "Khảo sát vừa được người tạo sửa — tải lại trang rồi trả lời lại")
    try:
        clean = svc.kiem_tra_tra_loi(questions, {a.question_id: a.value for a in body.answers})
    except ValueError as e:
        raise HTTPException(400, str(e))

    now = svc.now_str()
    try:
        if mine:
            rid = mine["id"]
            db.execute("UPDATE survey_responses SET updated_at = ? WHERE id = ?", (now, rid))
            db.execute("DELETE FROM survey_answers WHERE response_id = ?", (rid,))
        else:
            rid = db.execute(
                "INSERT INTO survey_responses (survey_id, staff_id, submitted_at) VALUES (?,?,?)",
                (survey_id, current["id"], now),
            ).lastrowid
        db.executemany(
            "INSERT INTO survey_answers (response_id, question_id, value) VALUES (?,?,?)",
            [(rid, qid, json.dumps(v, ensure_ascii=False)) for qid, v in clean.items()],
        )
    except sqlite3.IntegrityError:
        # Hai nguyên nhân, phải tra lại mới biết là cái nào — đừng báo "đã ghi nhận" khi
        # chưa có bài nào được lưu:
        # - Bấm Gửi hai lần liền: lần sau vỡ UNIQUE (survey_id, staff_id), bài lần đầu còn.
        # - Chủ vừa sửa / xoá khảo sát giữa lúc đọc câu hỏi và lúc ghi: vỡ FK, không có bài.
        db.rollback()
        if db.execute("SELECT 1 FROM survey_responses WHERE survey_id = ? AND staff_id = ?",
                      (survey_id, current["id"])).fetchone():
            raise HTTPException(409, "Câu trả lời của bạn đã được ghi nhận — tải lại trang để xem")
        raise HTTPException(409, "Khảo sát vừa được người tạo thay đổi — tải lại trang rồi trả lời lại")
    # Nhật ký chỉ ghi "đã nộp", KHÔNG ghi nội dung — khảo sát ẩn danh mà nội dung
    # nằm trong audit_logs thì ai xem được nhật ký là đọc được ai viết gì.
    write_audit(db, current["id"], "survey.edit_response" if mine else "survey.submit",
                "survey", survey_id, s["title"])
    db.commit()
    return {"ok": True, "updated": mine is not None}


# ── Kết quả ───────────────────────────────────────────────────────────────────
def _lay_de_xem_ket_qua(db, current, survey_id: int) -> dict:
    s = _lay(db, survey_id)
    if not _duoc_xem_ket_qua(db, current, s):
        raise HTTPException(403, "Không có quyền xem kết quả khảo sát này")
    s["state"] = svc.trang_thai(s)
    s["is_anonymous"] = bool(s["is_anonymous"])
    return s


@router.get("/{survey_id}/results")
def survey_results(
    survey_id: int,
    current: dict = Depends(require_feature("menu.surveys")),
    db: sqlite3.Connection = Depends(get_db),
):
    return svc.tong_hop(db, _lay_de_xem_ket_qua(db, current, survey_id))


_HEAD_FILL = PatternFill("solid", fgColor="FDE2E2")


def _ghi_bang(ws, header: list[str], rows: list[list], widths: Optional[list[int]] = None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = _HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _xlsx(res: dict) -> bytes:
    s = res["survey"]
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Tổng hợp"
    tong = []
    for st in res["stats"]:
        tong.append([st["title"], "", "", f"{st['answered']}/{st['total']} người trả lời"])
        for o in st.get("options", []):
            tong.append(["", o["label"], o["count"], f"{o['pct']}%"])
        if st.get("average") is not None:
            tong.append(["", "Điểm trung bình", st["average"], ""])
        for t in st.get("texts", []):
            tong.append(["", t["text"], "", t["staff_name"] or ""])
    _ghi_bang(ws, ["Câu hỏi", "Lựa chọn / Câu trả lời", "Số người", "Tỷ lệ"], tong, [50, 50, 12, 24])

    ws2 = wb.create_sheet("Câu trả lời")
    head = ([] if s["is_anonymous"] else ["Người trả lời", "Phòng", "Thời điểm nộp"]) + \
        [q["title"] for q in res["questions"]]
    body = [([] if s["is_anonymous"] else [r["staff_name"], r["dept_name"], r["submitted_at"]])
            + r["answers"] for r in res["rows"]]
    _ghi_bang(ws2, head, body)

    ws3 = wb.create_sheet("Tiến độ")
    _ghi_bang(ws3, ["Họ tên", "Phòng", "Trạng thái", "Thời điểm nộp"],
              [[r["staff_name"], r["dept_name"], "Đã trả lời" if r["responded"] else "Chưa trả lời",
                r["submitted_at"] or ""] for r in res["recipients"]], [30, 30, 16, 20])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/{survey_id}/export")
def export_results(
    survey_id: int,
    current: dict = Depends(require_feature("menu.surveys")),
    db: sqlite3.Connection = Depends(get_db),
):
    res = svc.tong_hop(db, _lay_de_xem_ket_qua(db, current, survey_id))
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in res["survey"]["title"])[:60]
    return Response(
        content=_xlsx(res),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers(f"Ket_qua_khao_sat_{safe.strip() or survey_id}.xlsx"),
    )
