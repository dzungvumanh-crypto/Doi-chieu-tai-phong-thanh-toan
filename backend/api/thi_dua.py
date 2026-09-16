"""Thi đua khen thưởng — Phòng Tổng hợp.

Ba loại dữ liệu độc lập, cùng một mức quyền phẳng (không có khái niệm "chủ sở
hữu bản ghi" như Khảo sát — ai có mã quyền quản lý thì sửa/xoá được mọi bản ghi,
giống `hr.edit_all`):

- `thi_dua.manage_unit`       — danh hiệu thi đua đơn vị (toàn Trung tâm / từng phòng)
- `thi_dua.manage_individual` — danh hiệu thi đua cá nhân theo cấp (Đảng / chuyên môn / công đoàn)
- `thi_dua.manage_initiative` — sáng kiến cá nhân được công nhận (kèm file quyết định)

`menu.thi_dua` là quyền xem/tra cứu chung cho cả ba loại. `thi_dua.export` gate
hai route xuất Excel.
"""
import io
import os
import re
import sqlite3
import unicodedata
from datetime import date as _date_cls, datetime as _datetime_cls, timedelta as _timedelta
from typing import Optional
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl.styles import Alignment, Font, PatternFill

from backend.core.concurrency import run_heavy
from backend.core.deps import require_feature
from backend.core.uploads import read_limited, safe_filename
from backend.database import _vn_now, get_db, write_audit
from backend.schemas.thi_dua import CaNhanIn, DonViIn, SangKienIn

router = APIRouter(prefix="/api/thi-dua", tags=["thi-dua"])

TOAN_TRUNG_TAM = "Toàn Trung tâm"

_TRAN_FILE = 15 * 1024 * 1024
_FILE_MO_RONG = {
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
}

_CAP_NHAN = {"dang": "Đảng", "chuyen_mon": "Chuyên môn", "cong_doan": "Công đoàn"}


# ── Tiện ích ──────────────────────────────────────────────────────────────────
def _download_headers(filename: str) -> dict:
    fallback = "".join(ch if ord(ch) < 128 and ch not in '\\"' else "_" for ch in filename)
    return {"Content-Disposition": (f'attachment; filename="{fallback}"; '
                                    f"filename*=UTF-8''{quote(filename, safe='')}")}


def _kieu_file(filename: str) -> tuple[str, str]:
    ten = safe_filename(filename, "quyet_dinh.dat")
    duoi = os.path.splitext(ten)[1].lower()
    if duoi not in _FILE_MO_RONG:
        raise HTTPException(
            400, f"File quyết định chỉ nhận {', '.join(sorted(_FILE_MO_RONG))} — "
                 f"file gửi lên là '{ten}'")
    return ten, _FILE_MO_RONG[duoi]


# ── Danh hiệu đơn vị ──────────────────────────────────────────────────────────
@router.get("/don-vi")
def list_don_vi(
    year: Optional[int] = Query(None),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    clauses, params = [], []
    if year is not None:
        clauses.append("t.year = ?")
        params.append(year)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    rows = db.execute(
        f"""SELECT t.id, t.year, t.department_id,
                   COALESCE(d.name, ?) AS department_name,
                   t.danh_hieu, t.so_quyet_dinh, t.ngay_quyet_dinh,
                   t.co_quan_ban_hanh, t.ghi_chu
            FROM thi_dua_don_vi t LEFT JOIN departments d ON d.id = t.department_id
            {where}
            ORDER BY t.year DESC, department_name, t.danh_hieu""",
        (TOAN_TRUNG_TAM, *params),
    ).fetchall()
    return [dict(r) for r in rows]


@router.post("/don-vi", status_code=201)
def create_don_vi(
    body: DonViIn,
    current: dict = Depends(require_feature("thi_dua.manage_unit")),
    db: sqlite3.Connection = Depends(get_db),
):
    if body.department_id is not None and not db.execute(
        "SELECT 1 FROM departments WHERE id = ?", (body.department_id,)
    ).fetchone():
        raise HTTPException(400, "Không tìm thấy phòng ban")
    now = _vn_now()
    cur = db.execute(
        """INSERT INTO thi_dua_don_vi
               (year, department_id, danh_hieu, so_quyet_dinh, ngay_quyet_dinh,
                co_quan_ban_hanh, ghi_chu, created_by, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (body.year, body.department_id, body.danh_hieu, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_ban_hanh, body.ghi_chu,
         current["id"], now, now),
    )
    write_audit(db, current["id"], "thi_dua.don_vi.create", "thi_dua_don_vi",
                cur.lastrowid, f"{body.year} — {body.danh_hieu}")
    db.commit()
    return {"id": cur.lastrowid}


@router.put("/don-vi/{item_id}")
def update_don_vi(
    item_id: int,
    body: DonViIn,
    current: dict = Depends(require_feature("thi_dua.manage_unit")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not db.execute("SELECT 1 FROM thi_dua_don_vi WHERE id = ?", (item_id,)).fetchone():
        raise HTTPException(404, "Không tìm thấy danh hiệu")
    if body.department_id is not None and not db.execute(
        "SELECT 1 FROM departments WHERE id = ?", (body.department_id,)
    ).fetchone():
        raise HTTPException(400, "Không tìm thấy phòng ban")
    db.execute(
        """UPDATE thi_dua_don_vi SET year=?, department_id=?, danh_hieu=?,
               so_quyet_dinh=?, ngay_quyet_dinh=?, co_quan_ban_hanh=?, ghi_chu=?,
               updated_at=?
           WHERE id=?""",
        (body.year, body.department_id, body.danh_hieu, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_ban_hanh, body.ghi_chu, _vn_now(), item_id),
    )
    write_audit(db, current["id"], "thi_dua.don_vi.update", "thi_dua_don_vi",
                item_id, f"{body.year} — {body.danh_hieu}")
    db.commit()
    return {"ok": True}


@router.delete("/don-vi/{item_id}")
def delete_don_vi(
    item_id: int,
    current: dict = Depends(require_feature("thi_dua.manage_unit")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not db.execute("SELECT 1 FROM thi_dua_don_vi WHERE id = ?", (item_id,)).fetchone():
        raise HTTPException(404, "Không tìm thấy danh hiệu")
    db.execute("DELETE FROM thi_dua_don_vi WHERE id = ?", (item_id,))
    write_audit(db, current["id"], "thi_dua.don_vi.delete", "thi_dua_don_vi", item_id)
    db.commit()
    return {"ok": True}


# ── Danh hiệu cá nhân ─────────────────────────────────────────────────────────
@router.get("/ca-nhan")
def list_ca_nhan(
    year: Optional[int] = Query(None),
    staff_id: Optional[int] = Query(None),
    cap: Optional[str] = Query(None),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    clauses, params = [], []
    if year is not None:
        clauses.append("t.year = ?")
        params.append(year)
    if staff_id is not None:
        clauses.append("t.staff_id = ?")
        params.append(staff_id)
    if cap is not None:
        clauses.append("t.cap = ?")
        params.append(cap)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    rows = db.execute(
        f"""SELECT t.id, t.staff_id, u.full_name AS staff_name, t.year, t.cap,
                   t.danh_hieu, t.so_quyet_dinh, t.ngay_quyet_dinh,
                   t.co_quan_ban_hanh, t.ghi_chu
            FROM thi_dua_ca_nhan t LEFT JOIN user_tttt u ON u.id = t.staff_id
            {where}
            ORDER BY t.year DESC, u.full_name""",
        params,
    ).fetchall()
    out = [dict(r) for r in rows]
    for r in out:
        r["cap_nhan"] = _CAP_NHAN.get(r["cap"], r["cap"])
    return out


def _staff_ton_tai(db, staff_id: int) -> bool:
    return bool(db.execute(
        "SELECT 1 FROM user_tttt WHERE id = ? AND (is_deleted = 0 OR is_deleted IS NULL)",
        (staff_id,),
    ).fetchone())


@router.post("/ca-nhan", status_code=201)
def create_ca_nhan(
    body: CaNhanIn,
    current: dict = Depends(require_feature("thi_dua.manage_individual")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not _staff_ton_tai(db, body.staff_id):
        raise HTTPException(400, "Không tìm thấy cán bộ")
    now = _vn_now()
    cur = db.execute(
        """INSERT INTO thi_dua_ca_nhan
               (staff_id, year, cap, danh_hieu, so_quyet_dinh, ngay_quyet_dinh,
                co_quan_ban_hanh, ghi_chu, created_by, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        (body.staff_id, body.year, body.cap.value, body.danh_hieu, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_ban_hanh, body.ghi_chu,
         current["id"], now, now),
    )
    write_audit(db, current["id"], "thi_dua.ca_nhan.create", "thi_dua_ca_nhan",
                cur.lastrowid, f"staff={body.staff_id}, {body.year} — {body.danh_hieu}")
    db.commit()
    return {"id": cur.lastrowid}


@router.put("/ca-nhan/{item_id}")
def update_ca_nhan(
    item_id: int,
    body: CaNhanIn,
    current: dict = Depends(require_feature("thi_dua.manage_individual")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not db.execute("SELECT 1 FROM thi_dua_ca_nhan WHERE id = ?", (item_id,)).fetchone():
        raise HTTPException(404, "Không tìm thấy danh hiệu")
    if not _staff_ton_tai(db, body.staff_id):
        raise HTTPException(400, "Không tìm thấy cán bộ")
    db.execute(
        """UPDATE thi_dua_ca_nhan SET staff_id=?, year=?, cap=?, danh_hieu=?,
               so_quyet_dinh=?, ngay_quyet_dinh=?, co_quan_ban_hanh=?, ghi_chu=?,
               updated_at=?
           WHERE id=?""",
        (body.staff_id, body.year, body.cap.value, body.danh_hieu, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_ban_hanh, body.ghi_chu, _vn_now(), item_id),
    )
    write_audit(db, current["id"], "thi_dua.ca_nhan.update", "thi_dua_ca_nhan",
                item_id, f"staff={body.staff_id}, {body.year} — {body.danh_hieu}")
    db.commit()
    return {"ok": True}


@router.delete("/ca-nhan/{item_id}")
def delete_ca_nhan(
    item_id: int,
    current: dict = Depends(require_feature("thi_dua.manage_individual")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not db.execute("SELECT 1 FROM thi_dua_ca_nhan WHERE id = ?", (item_id,)).fetchone():
        raise HTTPException(404, "Không tìm thấy danh hiệu")
    db.execute("DELETE FROM thi_dua_ca_nhan WHERE id = ?", (item_id,))
    write_audit(db, current["id"], "thi_dua.ca_nhan.delete", "thi_dua_ca_nhan", item_id)
    db.commit()
    return {"ok": True}


# ── Sáng kiến cá nhân ─────────────────────────────────────────────────────────
def _sang_kien_row(db, item_id: int) -> sqlite3.Row:
    row = db.execute(
        """SELECT sk.*, u.full_name AS staff_name
           FROM thi_dua_sang_kien sk LEFT JOIN user_tttt u ON u.id = sk.staff_id
           WHERE sk.id = ?""",
        (item_id,),
    ).fetchone()
    if not row:
        raise HTTPException(404, "Không tìm thấy sáng kiến")
    return row


def _sang_kien_out(row: sqlite3.Row) -> dict:
    d = dict(row)
    d["has_file"] = bool(d.pop("file_content", None))
    return d


@router.get("/sang-kien")
def list_sang_kien(
    year: Optional[int] = Query(None),
    staff_id: Optional[int] = Query(None),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    clauses, params = [], []
    if year is not None:
        clauses.append("sk.year = ?")
        params.append(year)
    if staff_id is not None:
        clauses.append("sk.staff_id = ?")
        params.append(staff_id)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    rows = db.execute(
        f"""SELECT sk.id, sk.staff_id, u.full_name AS staff_name, sk.year,
                   sk.ten_sang_kien, sk.so_quyet_dinh, sk.ngay_quyet_dinh,
                   sk.co_quan_cong_nhan, sk.ghi_chu, sk.file_name, sk.file_mime,
                   sk.file_size,
                   (sk.file_content IS NOT NULL) AS has_file
            FROM thi_dua_sang_kien sk LEFT JOIN user_tttt u ON u.id = sk.staff_id
            {where}
            ORDER BY sk.year DESC, u.full_name""",
        params,
    ).fetchall()
    return [dict(r) for r in rows]


@router.post("/sang-kien", status_code=201)
def create_sang_kien(
    body: SangKienIn,
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not _staff_ton_tai(db, body.staff_id):
        raise HTTPException(400, "Không tìm thấy cán bộ")
    now = _vn_now()
    cur = db.execute(
        """INSERT INTO thi_dua_sang_kien
               (staff_id, year, ten_sang_kien, so_quyet_dinh, ngay_quyet_dinh,
                co_quan_cong_nhan, ghi_chu, created_by, created_at, updated_at)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (body.staff_id, body.year, body.ten_sang_kien, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_cong_nhan, body.ghi_chu,
         current["id"], now, now),
    )
    write_audit(db, current["id"], "thi_dua.sang_kien.create", "thi_dua_sang_kien",
                cur.lastrowid, f"staff={body.staff_id}, {body.year} — {body.ten_sang_kien}")
    db.commit()
    return {"id": cur.lastrowid}


@router.put("/sang-kien/{item_id}")
def update_sang_kien(
    item_id: int,
    body: SangKienIn,
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    _sang_kien_row(db, item_id)
    if not _staff_ton_tai(db, body.staff_id):
        raise HTTPException(400, "Không tìm thấy cán bộ")
    db.execute(
        """UPDATE thi_dua_sang_kien SET staff_id=?, year=?, ten_sang_kien=?,
               so_quyet_dinh=?, ngay_quyet_dinh=?, co_quan_cong_nhan=?, ghi_chu=?,
               updated_at=?
           WHERE id=?""",
        (body.staff_id, body.year, body.ten_sang_kien, body.so_quyet_dinh,
         body.ngay_quyet_dinh, body.co_quan_cong_nhan, body.ghi_chu, _vn_now(), item_id),
    )
    write_audit(db, current["id"], "thi_dua.sang_kien.update", "thi_dua_sang_kien",
                item_id, f"staff={body.staff_id}, {body.year} — {body.ten_sang_kien}")
    db.commit()
    return {"ok": True}


@router.delete("/sang-kien/{item_id}")
def delete_sang_kien(
    item_id: int,
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    _sang_kien_row(db, item_id)
    db.execute("DELETE FROM thi_dua_sang_kien WHERE id = ?", (item_id,))
    write_audit(db, current["id"], "thi_dua.sang_kien.delete", "thi_dua_sang_kien", item_id)
    db.commit()
    return {"ok": True}


@router.post("/sang-kien/{item_id}/file")
async def upload_sang_kien_file(
    item_id: int,
    file: UploadFile = File(...),
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    _sang_kien_row(db, item_id)
    ten, mime = _kieu_file(file.filename)
    noi_dung = await read_limited(file, _TRAN_FILE, ten="File quyết định")
    db.execute(
        """UPDATE thi_dua_sang_kien
               SET file_name=?, file_mime=?, file_size=?, file_content=?, updated_at=?
           WHERE id=?""",
        (ten, mime, len(noi_dung), noi_dung, _vn_now(), item_id),
    )
    write_audit(db, current["id"], "thi_dua.sang_kien.upload_file", "thi_dua_sang_kien",
                item_id, ten)
    db.commit()
    return {"file_name": ten, "file_mime": mime, "file_size": len(noi_dung)}


@router.get("/sang-kien/{item_id}/file")
def download_sang_kien_file(
    item_id: int,
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    row = _sang_kien_row(db, item_id)
    if not row["file_content"]:
        raise HTTPException(404, "Sáng kiến này chưa có file quyết định")
    return Response(content=row["file_content"], media_type=row["file_mime"] or "application/pdf",
                    headers=_download_headers(row["file_name"] or "quyet_dinh.dat"))


@router.delete("/sang-kien/{item_id}/file")
def delete_sang_kien_file(
    item_id: int,
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    _sang_kien_row(db, item_id)
    db.execute(
        """UPDATE thi_dua_sang_kien
               SET file_name=NULL, file_mime=NULL, file_size=NULL, file_content=NULL, updated_at=?
           WHERE id=?""",
        (_vn_now(), item_id),
    )
    write_audit(db, current["id"], "thi_dua.sang_kien.delete_file", "thi_dua_sang_kien", item_id)
    db.commit()
    return {"ok": True}


# ── Nhập lô từ Excel ──────────────────────────────────────────────────────────
# Đơn vị có thể đang theo dõi thủ công bằng Excel trước khi có hệ thống — cần
# đường nhập lô, không chỉ gõ tay từng dòng. Không đoán khuôn Excel có sẵn của
# người dùng: hệ thống tự sinh file mẫu, dò cột theo TÊN tiêu đề (không theo vị
# trí cột cố định) để người dùng có thể thêm/bớt/đổi thứ tự cột tuỳ ý miễn còn
# đúng tên. `dry_run=True` (mặc định) chỉ xem trước, không ghi DB.
def _fold(s) -> str:
    """Hạ chữ + bỏ dấu + gộp khoảng trắng để dò tên cột bất kể cách gõ."""
    t = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    t = "".join(ch for ch in t if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", t).strip().lower()


def _doc_ngay(v) -> Optional[str]:
    """Excel trả về datetime, date, chuỗi dd/mm/yyyy hoặc số serial — chuẩn hoá về ISO."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, _datetime_cls):
        d = v.date()
    elif isinstance(v, _date_cls):
        d = v
    elif isinstance(v, (int, float)):
        try:
            d = _date_cls(1899, 12, 30) + _timedelta(days=int(v))
        except Exception:
            return None
    else:
        s = str(v).strip().split()[0]
        d = None
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y"):
            try:
                d = _datetime_cls.strptime(s, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None
    if not (2000 <= d.year <= 2100):
        return None
    return d.isoformat()


def _map_cap(v) -> Optional[str]:
    f = _fold(v)
    if not f:
        return None
    if f.startswith("dang"):
        return "dang"
    if "chuyen mon" in f:
        return "chuyen_mon"
    if "cong doan" in f:
        return "cong_doan"
    return None


def _dinh_vi_cot(ws, can_tim: dict, bat_buoc: set, max_row: int = 10):
    """Dò dòng tiêu đề trong `max_row` dòng đầu (khuôn `staff._parse_join_date_workbook`).

    `can_tim`: {trường: [các cụm từ đã fold, khớp kiểu "chứa trong ô"]}.
    Trả (header_row, {trường: chỉ_số_cột}) hoặc (None, {}) nếu thiếu cột bắt buộc.
    """
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, values_only=True), 1):
        found = {}
        for ci, cell in enumerate(row):
            h = _fold(cell)
            if not h:
                continue
            for truong, cum_tu in can_tim.items():
                if truong in found:
                    continue
                if any(c in h for c in cum_tu):
                    found[truong] = ci
        if bat_buoc <= set(found):
            return ri, found
    return None, {}


def _mau_workbook(cols: list[tuple[str, int]], vi_du: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mẫu nhập"
    fill = PatternFill("solid", fgColor="FEE2E2")
    for j, (ten, w) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=ten)
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    for j, v in enumerate(vi_du, start=1):
        ws.cell(row=2, column=j, value=v)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# — Đơn vị —
_CAN_DON_VI = {
    "nam": ["nam"], "don_vi": ["don vi", "phong"], "danh_hieu": ["danh hieu"],
    "so_qd": ["so quyet dinh"], "ngay_qd": ["ngay quyet dinh"],
    "co_quan": ["co quan ban hanh"], "ghi_chu": ["ghi chu"],
}
_BAT_BUOC_DON_VI = {"nam", "danh_hieu"}
_TOAN_TTAM_ALIAS = {"", "toan trung tam", "trung tam", "toan bo", "ca trung tam"}


def _doc_wb_don_vi(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    header_row, cot = _dinh_vi_cot(ws, _CAN_DON_VI, _BAT_BUOC_DON_VI)
    if header_row is None:
        wb.close()
        raise HTTPException(
            400, "Không tìm thấy đủ cột bắt buộc (Năm, Danh hiệu) trong file. "
                 "Hãy tải file mẫu và điền đúng tên cột.")
    items = []
    for ri, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        def g(key):
            ci = cot.get(key)
            return row[ci] if ci is not None and ci < len(row) else None
        nam, danh_hieu = g("nam"), g("danh_hieu")
        if not str(nam or "").strip() and not str(danh_hieu or "").strip():
            continue
        items.append({"dong": ri, "nam": nam, "don_vi": g("don_vi"), "danh_hieu": danh_hieu,
                      "so_qd": g("so_qd"), "ngay_qd": g("ngay_qd"), "co_quan": g("co_quan"),
                      "ghi_chu": g("ghi_chu")})
    wb.close()
    return items


@router.get("/don-vi/import-template")
def don_vi_import_template(current: dict = Depends(require_feature("thi_dua.manage_unit"))):
    noi_dung = _mau_workbook(
        [("Năm", 10), ("Đơn vị (để trống = Toàn Trung tâm)", 32), ("Danh hiệu", 34),
         ("Số quyết định", 18), ("Ngày quyết định", 16), ("Cơ quan ban hành", 28), ("Ghi chú", 30)],
        [2026, "", "Tập thể lao động xuất sắc", "123/QĐ-TTTT", "15/01/2026", "Agribank", ""],
    )
    return Response(content=noi_dung,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers("mau_danh_hieu_don_vi.xlsx"))


@router.post("/don-vi/import")
async def import_don_vi(
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    current: dict = Depends(require_feature("thi_dua.manage_unit")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Chỉ nhận file Excel .xlsx/.xlsm")
    noi_dung = await read_limited(file, ten="File Excel danh hiệu đơn vị")
    items = await run_heavy(_doc_wb_don_vi, noi_dung)
    if not items:
        raise HTTPException(400, "Không đọc được dòng dữ liệu nào trong file")

    depts = {_fold(d["name"]): d["id"] for d in db.execute("SELECT id, name FROM departments")}
    loi, them = [], 0
    now = _vn_now()
    for it in items:
        try:
            nam_i = int(it["nam"])
        except (TypeError, ValueError):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: '{it['nam']}'"})
            continue
        if not (2000 <= nam_i <= 2100):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: {nam_i}"})
            continue
        danh_hieu = str(it["danh_hieu"] or "").strip()
        if not danh_hieu:
            loi.append({"dong": it["dong"], "ly_do": "Thiếu tên danh hiệu"})
            continue
        dv_fold = _fold(it["don_vi"])
        if dv_fold in _TOAN_TTAM_ALIAS:
            dept_id = None
        else:
            dept_id = depts.get(dv_fold)
            if dept_id is None:
                loi.append({"dong": it["dong"], "ly_do": f"Không khớp tên đơn vị: '{it['don_vi']}'"})
                continue
        ngay_qd = _doc_ngay(it["ngay_qd"])
        if it["ngay_qd"] and not ngay_qd:
            loi.append({"dong": it["dong"], "ly_do": f"Ngày quyết định không đọc được: '{it['ngay_qd']}'"})
            continue
        if not dry_run:
            db.execute(
                """INSERT INTO thi_dua_don_vi
                       (year, department_id, danh_hieu, so_quyet_dinh, ngay_quyet_dinh,
                        co_quan_ban_hanh, ghi_chu, created_by, created_at, updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (nam_i, dept_id, danh_hieu, str(it["so_qd"] or "").strip() or None, ngay_qd,
                 str(it["co_quan"] or "").strip() or None, str(it["ghi_chu"] or "").strip() or None,
                 current["id"], now, now),
            )
        them += 1

    if not dry_run and them:
        write_audit(db, current["id"], "thi_dua.don_vi.import", "thi_dua_don_vi", None,
                    f"{file.filename}: thêm {them}, lỗi {len(loi)}")
        db.commit()
    return {"tong_dong": len(items), "da_them": them, "loi": loi}


# — Cá nhân —
_CAN_CA_NHAN = {
    "ma_cb": ["ma can bo", "ma cb", "ma nhan vien"], "cap": ["cap"],
    "danh_hieu": ["danh hieu"], "nam": ["nam"], "so_qd": ["so quyet dinh"],
    "ngay_qd": ["ngay quyet dinh"], "co_quan": ["co quan ban hanh"], "ghi_chu": ["ghi chu"],
}
_BAT_BUOC_CA_NHAN = {"ma_cb", "cap", "danh_hieu", "nam"}


def _doc_wb_ca_nhan(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    header_row, cot = _dinh_vi_cot(ws, _CAN_CA_NHAN, _BAT_BUOC_CA_NHAN)
    if header_row is None:
        wb.close()
        raise HTTPException(
            400, "Không tìm thấy đủ cột bắt buộc (Mã cán bộ, Cấp, Danh hiệu, Năm) trong file. "
                 "Hãy tải file mẫu và điền đúng tên cột.")
    items = []
    for ri, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        def g(key):
            ci = cot.get(key)
            return row[ci] if ci is not None and ci < len(row) else None
        ma_cb = g("ma_cb")
        if not str(ma_cb or "").strip():
            continue
        if isinstance(ma_cb, float) and ma_cb.is_integer():
            ma_cb = int(ma_cb)
        items.append({"dong": ri, "ma_cb": ma_cb, "cap": g("cap"), "danh_hieu": g("danh_hieu"),
                      "nam": g("nam"), "so_qd": g("so_qd"), "ngay_qd": g("ngay_qd"),
                      "co_quan": g("co_quan"), "ghi_chu": g("ghi_chu")})
    wb.close()
    return items


@router.get("/ca-nhan/import-template")
def ca_nhan_import_template(current: dict = Depends(require_feature("thi_dua.manage_individual"))):
    noi_dung = _mau_workbook(
        [("Mã cán bộ", 14), ("Họ và tên (chỉ để đối chiếu)", 26), ("Năm", 10),
         ("Cấp (Đảng/Chuyên môn/Công đoàn)", 26), ("Danh hiệu", 34), ("Số quyết định", 18),
         ("Ngày quyết định", 16), ("Cơ quan ban hành", 28), ("Ghi chú", 30)],
        ["NV001", "Nguyễn Văn A", 2026, "Chuyên môn", "Lao động tiên tiến",
         "45/QĐ-TTTT", "20/01/2026", "Agribank", ""],
    )
    return Response(content=noi_dung,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers("mau_danh_hieu_ca_nhan.xlsx"))


@router.post("/ca-nhan/import")
async def import_ca_nhan(
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    current: dict = Depends(require_feature("thi_dua.manage_individual")),
    db: sqlite3.Connection = Depends(get_db),
):
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Chỉ nhận file Excel .xlsx/.xlsm")
    noi_dung = await read_limited(file, ten="File Excel danh hiệu cá nhân")
    items = await run_heavy(_doc_wb_ca_nhan, noi_dung)
    if not items:
        raise HTTPException(400, "Không đọc được dòng dữ liệu nào trong file")

    staff_by_code = {
        str(r["employee_code"] or "").strip(): r["id"]
        for r in db.execute(
            "SELECT id, employee_code FROM user_tttt WHERE is_deleted = 0 OR is_deleted IS NULL")
    }
    loi, them = [], 0
    now = _vn_now()
    for it in items:
        ma_cb = str(it["ma_cb"]).strip()
        staff_id = staff_by_code.get(ma_cb)
        if staff_id is None:
            loi.append({"dong": it["dong"], "ly_do": f"Không khớp mã cán bộ: '{ma_cb}'"})
            continue
        cap = _map_cap(it["cap"])
        if cap is None:
            loi.append({"dong": it["dong"],
                        "ly_do": f"Cấp không hợp lệ: '{it['cap']}' (cần Đảng/Chuyên môn/Công đoàn)"})
            continue
        danh_hieu = str(it["danh_hieu"] or "").strip()
        if not danh_hieu:
            loi.append({"dong": it["dong"], "ly_do": "Thiếu tên danh hiệu"})
            continue
        try:
            nam_i = int(it["nam"])
        except (TypeError, ValueError):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: '{it['nam']}'"})
            continue
        if not (2000 <= nam_i <= 2100):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: {nam_i}"})
            continue
        ngay_qd = _doc_ngay(it["ngay_qd"])
        if it["ngay_qd"] and not ngay_qd:
            loi.append({"dong": it["dong"], "ly_do": f"Ngày quyết định không đọc được: '{it['ngay_qd']}'"})
            continue
        if not dry_run:
            db.execute(
                """INSERT INTO thi_dua_ca_nhan
                       (staff_id, year, cap, danh_hieu, so_quyet_dinh, ngay_quyet_dinh,
                        co_quan_ban_hanh, ghi_chu, created_by, created_at, updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (staff_id, nam_i, cap, danh_hieu, str(it["so_qd"] or "").strip() or None, ngay_qd,
                 str(it["co_quan"] or "").strip() or None, str(it["ghi_chu"] or "").strip() or None,
                 current["id"], now, now),
            )
        them += 1

    if not dry_run and them:
        write_audit(db, current["id"], "thi_dua.ca_nhan.import", "thi_dua_ca_nhan", None,
                    f"{file.filename}: thêm {them}, lỗi {len(loi)}")
        db.commit()
    return {"tong_dong": len(items), "da_them": them, "loi": loi}


# — Sáng kiến —
_CAN_SANG_KIEN = {
    "ma_cb": ["ma can bo", "ma cb", "ma nhan vien"], "ten_sk": ["ten sang kien"],
    "nam": ["nam"], "so_qd": ["so quyet dinh"], "ngay_qd": ["ngay quyet dinh"],
    "co_quan": ["co quan cong nhan"], "ghi_chu": ["ghi chu"],
}
_BAT_BUOC_SANG_KIEN = {"ma_cb", "ten_sk", "nam"}


def _doc_wb_sang_kien(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    header_row, cot = _dinh_vi_cot(ws, _CAN_SANG_KIEN, _BAT_BUOC_SANG_KIEN)
    if header_row is None:
        wb.close()
        raise HTTPException(
            400, "Không tìm thấy đủ cột bắt buộc (Mã cán bộ, Tên sáng kiến, Năm) trong file. "
                 "Hãy tải file mẫu và điền đúng tên cột.")
    items = []
    for ri, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        def g(key):
            ci = cot.get(key)
            return row[ci] if ci is not None and ci < len(row) else None
        ma_cb = g("ma_cb")
        if not str(ma_cb or "").strip():
            continue
        if isinstance(ma_cb, float) and ma_cb.is_integer():
            ma_cb = int(ma_cb)
        items.append({"dong": ri, "ma_cb": ma_cb, "ten_sk": g("ten_sk"), "nam": g("nam"),
                      "so_qd": g("so_qd"), "ngay_qd": g("ngay_qd"), "co_quan": g("co_quan"),
                      "ghi_chu": g("ghi_chu")})
    wb.close()
    return items


@router.get("/sang-kien/import-template")
def sang_kien_import_template(current: dict = Depends(require_feature("thi_dua.manage_initiative"))):
    noi_dung = _mau_workbook(
        [("Mã cán bộ", 14), ("Họ và tên (chỉ để đối chiếu)", 26), ("Năm", 10),
         ("Tên sáng kiến", 36), ("Số quyết định", 18), ("Ngày quyết định", 16),
         ("Cơ quan công nhận", 28), ("Ghi chú", 30)],
        ["NV001", "Nguyễn Văn A", 2026, "Giải pháp cải tiến quy trình đối chiếu",
         "78/QĐ-TTTT", "10/03/2026", "Agribank", "File quyết định đính kèm sau khi nhập"],
    )
    return Response(content=noi_dung,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers("mau_sang_kien_ca_nhan.xlsx"))


@router.post("/sang-kien/import")
async def import_sang_kien(
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    current: dict = Depends(require_feature("thi_dua.manage_initiative")),
    db: sqlite3.Connection = Depends(get_db),
):
    """Nhập lô sáng kiến — KHÔNG kèm file quyết định (Excel không chở được file);
    gắn file quyết định từng dòng sau khi nhập, qua nút Sửa."""
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Chỉ nhận file Excel .xlsx/.xlsm")
    noi_dung = await read_limited(file, ten="File Excel sáng kiến")
    items = await run_heavy(_doc_wb_sang_kien, noi_dung)
    if not items:
        raise HTTPException(400, "Không đọc được dòng dữ liệu nào trong file")

    staff_by_code = {
        str(r["employee_code"] or "").strip(): r["id"]
        for r in db.execute(
            "SELECT id, employee_code FROM user_tttt WHERE is_deleted = 0 OR is_deleted IS NULL")
    }
    loi, them = [], 0
    now = _vn_now()
    for it in items:
        ma_cb = str(it["ma_cb"]).strip()
        staff_id = staff_by_code.get(ma_cb)
        if staff_id is None:
            loi.append({"dong": it["dong"], "ly_do": f"Không khớp mã cán bộ: '{ma_cb}'"})
            continue
        ten_sk = str(it["ten_sk"] or "").strip()
        if not ten_sk:
            loi.append({"dong": it["dong"], "ly_do": "Thiếu tên sáng kiến"})
            continue
        try:
            nam_i = int(it["nam"])
        except (TypeError, ValueError):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: '{it['nam']}'"})
            continue
        if not (2000 <= nam_i <= 2100):
            loi.append({"dong": it["dong"], "ly_do": f"Năm không hợp lệ: {nam_i}"})
            continue
        ngay_qd = _doc_ngay(it["ngay_qd"])
        if it["ngay_qd"] and not ngay_qd:
            loi.append({"dong": it["dong"], "ly_do": f"Ngày quyết định không đọc được: '{it['ngay_qd']}'"})
            continue
        if not dry_run:
            db.execute(
                """INSERT INTO thi_dua_sang_kien
                       (staff_id, year, ten_sang_kien, so_quyet_dinh, ngay_quyet_dinh,
                        co_quan_cong_nhan, ghi_chu, created_by, created_at, updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (staff_id, nam_i, ten_sk, str(it["so_qd"] or "").strip() or None, ngay_qd,
                 str(it["co_quan"] or "").strip() or None, str(it["ghi_chu"] or "").strip() or None,
                 current["id"], now, now),
            )
        them += 1

    if not dry_run and them:
        write_audit(db, current["id"], "thi_dua.sang_kien.import", "thi_dua_sang_kien", None,
                    f"{file.filename}: thêm {them}, lỗi {len(loi)}")
        db.commit()
    return {"tong_dong": len(items), "da_them": them, "loi": loi}


# ── Gợi ý tên danh hiệu (autocomplete) ────────────────────────────────────────
@router.get("/danh-hieu-goi-y")
def goi_y_danh_hieu(
    loai: str = Query(..., pattern="^(don_vi|ca_nhan)$"),
    q: str = Query(""),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    # Không lọc gắt theo `q` ở server: frontend dùng danh sách này làm
    # `autocomplete=[...]` gốc cho ui.input (lọc theo tiền tố ngay trên trình
    # duyệt), nên cần trả đủ rộng ngay từ lần gọi đầu, không phải gọi lại theo
    # từng phím gõ.
    table = "thi_dua_don_vi" if loai == "don_vi" else "thi_dua_ca_nhan"
    rows = db.execute(
        f"""SELECT DISTINCT danh_hieu FROM {table}
            WHERE danh_hieu LIKE ? ORDER BY danh_hieu LIMIT 300""",
        (f"%{q.strip()}%",),
    ).fetchall()
    return [r["danh_hieu"] for r in rows]


# ── Tra cứu, thống kê ─────────────────────────────────────────────────────────
def _tong_hop_rows(db: sqlite3.Connection, year: Optional[int]) -> list[dict]:
    clause = "WHERE t.year = ?" if year is not None else ""
    params = (year,) if year is not None else ()
    don_vi = db.execute(
        f"""SELECT 'Đơn vị' AS loai, t.year, COALESCE(d.name, ?) AS doi_tuong,
                   NULL AS cap_nhan, t.danh_hieu, t.so_quyet_dinh, t.ngay_quyet_dinh,
                   t.co_quan_ban_hanh
            FROM thi_dua_don_vi t LEFT JOIN departments d ON d.id = t.department_id
            {clause}""",
        (TOAN_TRUNG_TAM, *params),
    ).fetchall()
    ca_nhan = db.execute(
        f"""SELECT 'Cá nhân' AS loai, t.year, u.full_name AS doi_tuong,
                   t.cap AS cap_nhan, t.danh_hieu, t.so_quyet_dinh, t.ngay_quyet_dinh,
                   t.co_quan_ban_hanh
            FROM thi_dua_ca_nhan t LEFT JOIN user_tttt u ON u.id = t.staff_id
            {clause}""",
        params,
    ).fetchall()
    out = [dict(r) for r in (*don_vi, *ca_nhan)]
    for r in out:
        if r["cap_nhan"]:
            r["cap_nhan"] = _CAP_NHAN.get(r["cap_nhan"], r["cap_nhan"])
    out.sort(key=lambda r: (-r["year"], r["loai"], r["doi_tuong"] or ""))
    return out


@router.get("/stats/tong-hop")
def stats_tong_hop(
    year: Optional[int] = Query(None),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    return _tong_hop_rows(db, year)


def _sang_kien_stats_rows(db: sqlite3.Connection, staff_id: Optional[int],
                          year: Optional[int]) -> list[dict]:
    clauses, params = [], []
    if staff_id is not None:
        clauses.append("sk.staff_id = ?")
        params.append(staff_id)
    if year is not None:
        clauses.append("sk.year = ?")
        params.append(year)
    where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
    rows = db.execute(
        f"""SELECT u.full_name AS doi_tuong, sk.year, sk.ten_sang_kien,
                   sk.so_quyet_dinh, sk.ngay_quyet_dinh, sk.co_quan_cong_nhan,
                   (sk.file_content IS NOT NULL) AS has_file
            FROM thi_dua_sang_kien sk LEFT JOIN user_tttt u ON u.id = sk.staff_id
            {where}
            ORDER BY u.full_name, sk.year DESC""",
        params,
    ).fetchall()
    return [dict(r) for r in rows]


@router.get("/stats/sang-kien")
def stats_sang_kien(
    staff_id: Optional[int] = Query(None),
    year: Optional[int] = Query(None),
    current: dict = Depends(require_feature("menu.thi_dua")),
    db: sqlite3.Connection = Depends(get_db),
):
    return _sang_kien_stats_rows(db, staff_id, year)


# ── Xuất Excel ────────────────────────────────────────────────────────────────
def _workbook(rows: list[dict], tieu_de: str, cot: list[tuple[str, str, int]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách"
    ws.cell(row=1, column=1, value=tieu_de).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cot) + 1)
    fill = PatternFill("solid", fgColor="FEE2E2")
    for j, nhan in enumerate(["STT"] + [n for _, n, _ in cot], start=1):
        c = ws.cell(row=3, column=j, value=nhan)
        c.font, c.fill = Font(bold=True), fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions["A"].width = 6
    for j, (_, _, w) in enumerate(cot, start=2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    for i, r in enumerate(rows, start=1):
        ws.cell(row=3 + i, column=1, value=i)
        for j, (f, _, _) in enumerate(cot, start=2):
            ws.cell(row=3 + i, column=j, value=r.get(f))
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_COT_TONG_HOP = [
    ("loai", "Loại", 10), ("year", "Năm", 8), ("doi_tuong", "Đơn vị / Cán bộ", 28),
    ("cap_nhan", "Cấp", 12), ("danh_hieu", "Danh hiệu", 32),
    ("so_quyet_dinh", "Số quyết định", 16), ("ngay_quyet_dinh", "Ngày quyết định", 14),
    ("co_quan_ban_hanh", "Cơ quan ban hành", 26),
]
_COT_SANG_KIEN = [
    ("doi_tuong", "Cán bộ", 26), ("year", "Năm", 8), ("ten_sang_kien", "Tên sáng kiến", 40),
    ("so_quyet_dinh", "Số quyết định", 16), ("ngay_quyet_dinh", "Ngày quyết định", 14),
    ("co_quan_cong_nhan", "Cơ quan công nhận", 26),
]


@router.get("/export/tong-hop")
async def export_tong_hop(
    year: Optional[int] = Query(None),
    current: dict = Depends(require_feature("thi_dua.export")),
    db: sqlite3.Connection = Depends(get_db),
):
    rows = _tong_hop_rows(db, year)
    tieu_de = "TỔNG HỢP THI ĐUA KHEN THƯỞNG" + (f" — NĂM {year}" if year else "")
    noi_dung = await run_heavy(_workbook, rows, tieu_de, _COT_TONG_HOP)
    write_audit(db, current["id"], "thi_dua.export_tong_hop", "thi_dua", None,
                f"năm={year}, {len(rows)} dòng")
    db.commit()
    ten_file = f"tong_hop_thi_dua_{year or 'tat_ca'}.xlsx"
    return Response(content=noi_dung,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers(ten_file))


@router.get("/export/sang-kien")
async def export_sang_kien(
    staff_id: Optional[int] = Query(None),
    year: Optional[int] = Query(None),
    current: dict = Depends(require_feature("thi_dua.export")),
    db: sqlite3.Connection = Depends(get_db),
):
    rows = _sang_kien_stats_rows(db, staff_id, year)
    noi_dung = await run_heavy(_workbook, rows, "DANH SÁCH SÁNG KIẾN THEO CÁ NHÂN", _COT_SANG_KIEN)
    write_audit(db, current["id"], "thi_dua.export_sang_kien", "thi_dua", None,
                f"staff_id={staff_id}, năm={year}, {len(rows)} dòng")
    db.commit()
    ten_file = f"sang_kien_ca_nhan_{year or 'tat_ca'}.xlsx"
    return Response(content=noi_dung,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_download_headers(ten_file))
