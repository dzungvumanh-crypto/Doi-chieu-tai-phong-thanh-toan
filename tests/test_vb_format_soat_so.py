"""Soát thứ tự đánh số — logic mượn từ legal-merger (máy trạng thái + che trích dẫn).

Chỉ BÁO, không sửa: test khẳng định cả chỗ phải báo lẫn chỗ phải im lặng
(Điều liên tục qua Chương, Quy chế kèm Quyết định đánh lại Điều 1, nội dung
trích dẫn của văn bản sửa đổi).
"""
import io

from docx import Document

from backend.services.vb_format import nhan_dien, soat_so
from backend.services.vb_format.chuan_hoa import chuan_hoa


def _soat(dong: list[str], trong_bang: list[bool] | None = None) -> list[dict]:
    tb = trong_bang or [False] * len(dong)
    ma = nhan_dien.phan_loai(list(zip(dong, tb)))
    return soat_so.soat_thu_tu(ma, dong, tb)


# ── Không báo khi đúng ───────────────────────────────────────────────────────
def test_van_ban_dung_khong_bao():
    assert _soat([
        "Điều 1. Phạm vi",
        "1. Khoản một.",
        "a) Điểm a;",
        "b) Điểm b.",
        "2. Khoản hai.",
        "Điều 2. Hiệu lực",
        "1. Khoản một của Điều 2.",
    ]) == []


def test_dieu_lien_tuc_qua_chuong_khong_bao():
    assert _soat([
        "Chương I", "QUY ĐỊNH CHUNG",
        "Điều 1. Phạm vi", "Điều 2. Đối tượng",
        "Chương II", "QUY ĐỊNH CỤ THỂ",
        "Điều 3. Nội dung",
    ]) == []


def test_quy_che_kem_quyet_dinh_danh_lai_dieu_1():
    assert _soat([
        "Điều 1. Ban hành kèm theo Quyết định này Quy chế.",
        "Điều 2. Hiệu lực.",
        "Điều 3. Trách nhiệm thi hành.",
        "QUY CHẾ",
        "Chương I",
        "QUY ĐỊNH CHUNG",
        "Điều 1. Phạm vi điều chỉnh",
    ]) == []


# ── Báo khi sai ──────────────────────────────────────────────────────────────
def test_nhay_dieu():
    kq = _soat(["Điều 1. A", "Điều 3. C"])
    assert len(kq) == 1
    assert kq[0]["stt"] == 2 and "thiếu Điều 2" in kq[0]["loi"]


def test_diem_bo_qua_chu_d_gach():
    kq = _soat(["Điều 1. A", "1. Khoản.", "a) x;", "b) x;", "c) x;", "d) x;", "e) x."])
    assert len(kq) == 1
    assert "thiếu điểm đ)" in kq[0]["loi"] and "12.5.b" in kq[0]["loi"]


def test_trung_so_khoan():
    kq = _soat(["Điều 1. A", "1. Một.", "2. Hai.", "2. Hai nữa."])
    assert len(kq) == 1 and "trùng số" in kq[0]["loi"]


def test_khoan_danh_lai_giua_dieu():
    kq = _soat(["Điều 1. A", "1. Một.", "2. Hai.", "1. Lại một."])
    assert len(kq) == 1 and "đánh lại từ đầu" in kq[0]["loi"]


def test_mot_cho_sai_khong_keo_ca_day():
    # Lấy số mới làm mốc: sau "Điều 3" thì "Điều 4" là đúng, không báo tiếp.
    kq = _soat(["Điều 1. A", "Điều 3. C", "Điều 4. D", "Điều 5. E"])
    assert [x["stt"] for x in kq] == [2]


def test_khoan_dau_tien_khong_phai_1_trong_dieu():
    kq = _soat(["Điều 1. A", "2. Hai."])
    assert len(kq) == 1 and "thiếu khoản 1." in kq[0]["loi"]


def test_muc_la_ma_nhay():
    kq = _soat(["I. TÌNH HÌNH", "II. KẾT QUẢ", "IV. KIẾN NGHỊ"])
    assert len(kq) == 1 and "thiếu mục III." in kq[0]["loi"]


# ── Tiểu khoản và tiết ───────────────────────────────────────────────────────
def test_tieu_khoan_sai_so_khoan_cha():
    kq = _soat(["Điều 1. A", "1. Khoản một", "1.1 Nội dung", "2.2 Nội dung"])
    assert len(kq) == 1 and "nằm dưới khoản 1" in kq[0]["loi"]


def test_tieu_khoan_nhay():
    kq = _soat(["Điều 1. A", "2. Khoản hai" , "2.1 Một", "2.3 Ba"])
    loi = [x["loi"] for x in kq]
    assert any("thiếu 2.2" in x for x in loi)


def test_tiet_la_ma_dai_truoc_ngan():
    # "(xi)" phải đọc là 11, không phải khớp "x" rồi trượt.
    dong = ["Điều 1. A", "1. K", "a) Đ"] + [f"({r}) t" for r in
            ("i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi")]
    assert _soat(dong) == []
    kq = _soat(["Điều 1. A", "1. K", "a) Đ", "(i) t", "(iii) t"])
    assert len(kq) == 1 and "thiếu tiết (ii)" in kq[0]["loi"]


def test_so_tien_khong_phai_tieu_khoan():
    assert _soat(["Điều 1. A", "1. K", "1.500 đồng là phí mỗi món."]) == []


# ── Che trích dẫn ────────────────────────────────────────────────────────────
def test_noi_dung_trich_dan_nhieu_doan_bi_che():
    assert _soat([
        "Điều 1. Sửa đổi Điều 5 như sau:",
        "“Điều 5. Nội dung mới",
        "1. Khoản một.",
        "3. Khoản ba của văn bản được sửa.",
        "a) Điểm a.”",
        "Điều 2. Hiệu lực.",
    ]) == []


def test_trich_dan_mot_doan_bi_che():
    assert _soat([
        "Điều 1. Sửa đổi",
        "1. Sửa khoản 3 như sau:",
        "“5. Nội dung khoản năm.”.",
        "2. Khoản hai.",
    ]) == []


def test_cau_mo_bang_ngoac_khong_che_doan_sau():
    # Câu thường mở bằng “…” và đóng ngay trong câu — không được che các đoạn
    # sau tới dấu đóng ngoặc kế tiếp.
    kq = _soat([
        "Điều 1. A",
        "“Chuyển đổi số” là nhiệm vụ trọng tâm.",
        "Điều 3. C",
        "Điều 4. D có cụm “x”.",
    ])
    assert [x["stt"] for x in kq] == [3]


def test_ngoac_mo_khong_dong_thi_khong_che():
    kq = _soat(["Điều 1. A", "“Trích dẫn không đóng", "Điều 3. C"])
    assert [x["stt"] for x in kq] == [3]


def test_quoc_hieu_trong_bang_mo_van_ban_moi():
    # Mẫu biểu ở phụ lục: Quốc hiệu riêng (trong bảng) rồi đánh "1." lại.
    dong = ["Điều 1. A", "1. x", "2. y", "3. z",
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "1. Mục của mẫu biểu"]
    assert _soat(dong, [False, False, False, False, True, False]) == []


def test_ten_chuong_co_chu_quyen_khong_thanh_chuc_danh():
    # Có phần đầu văn bản thật: thiếu tên loại thì dòng in hoa ở đầu bị luật
    # "tên đơn vị" nhận trước — test đo nhầm thứ khác.
    dong = ["Số: 15/2024/TT-NHNN", "THÔNG TƯ", "Quy định về dịch vụ thanh toán",
            "Chương I", "QUY ĐỊNH CHUNG", "Điều 1. A",
            "Chương II", "QUYỀN VÀ TRÁCH NHIỆM", "Điều 2. B"]
    ma = nhan_dien.phan_loai([(d, False) for d in dong])
    assert ma[7] == "tieu_de_phan_chuong"
    assert _soat(dong) == []


def test_style_tieu_de_tach_danh_sach_buoc():
    # Tài liệu hướng dẫn: "1." Heading 1 → "1.1" Heading 2 → các bước "1. 2." thường.
    dong = ["1. Truy cập", "1.1. Đăng nhập", "1. Mở trình duyệt.", "2. Nhập địa chỉ.",
            "1.2. Điều hướng", "1. Trỏ chuột.", "2. Truy cập", "2.1. Tổng quan", "1. Bước."]
    lvl = [1, 2, None, None, 2, None, 1, 2, None]
    tb = [False] * len(dong)
    ma = nhan_dien.phan_loai(list(zip(dong, tb)))
    assert soat_so.soat_thu_tu(ma, dong, tb, lvl) == []
    # Không có style thì chính các dòng đó bị báo — đúng lý do phải đọc style.
    assert soat_so.soat_thu_tu(ma, dong, tb) != []


def test_dieu_dat_style_khong_dong_nhat_van_mot_day():
    # TT 15/2024: Điều 8 có outline level, Điều 7 không — ở đây Điều 2.
    dong = ["Điều 1. A", "1. x", "2. y", "Điều 2. B", "1. x", "Điều 3. C"]
    lvl = [None, None, None, 3, None, None]
    tb = [False] * len(dong)
    ma = nhan_dien.phan_loai(list(zip(dong, tb)))
    assert soat_so.soat_thu_tu(ma, dong, tb, lvl) == []


def test_van_ban_co_dieu_bo_qua_style_tieu_de():
    # Khoản 2 lệch outline level, tiêu đề không số chen giữa khoản 2 và 3.
    dong = ["Điều 1. A", "1. x", "2. y", "Bảng tổng hợp", "3. z", "Điều 2. B"]
    lvl = [None, None, 3, 2, None, None]
    tb = [False] * len(dong)
    ma = nhan_dien.phan_loai(list(zip(dong, tb)))
    assert soat_so.soat_thu_tu(ma, dong, tb, lvl) == []


def test_so_rat_lon_khong_dung_ca_day_thieu():
    kq = _soat(["Điều 1. A", "Điều 300000000. B"])
    assert len(kq) == 1 and kq[0]["loi"].endswith("Điều 2, Điều 3, Điều 4 …")


def test_style_tu_tham_chieu_khong_treo():
    doc = Document()
    st = doc.styles.add_style("Vong", 1)
    st.element.get_or_add_basedOn().val = st.style_id
    p = doc.add_paragraph("1. x", style="Vong")
    assert soat_so.muc_de_muc(p) is None


def test_o_bang_bo_qua():
    kq = _soat(["Điều 1. A", "1. x", "5. y"], [False, True, True])
    assert kq == []


# ── Qua chuan_hoa: có trong báo cáo, tắt được ────────────────────────────────
def _docx(dong: list[str]) -> bytes:
    doc = Document()
    for d in dong:
        doc.add_paragraph(d)
    ra = io.BytesIO()
    doc.save(ra)
    return ra.getvalue()


def test_chuan_hoa_tra_soat_so_va_tat_duoc():
    du_lieu = _docx(["Điều 1. A", "Điều 3. C"])
    _, bc = chuan_hoa(du_lieu)
    assert len(bc["soat_so"]) == 1 and bc["soat_so"][0]["stt"] == 2

    _, bc = chuan_hoa(du_lieu, {"danh_so": {"soat_thu_tu": False}})
    assert bc["soat_so"] == []


def test_chuan_hoa_khong_sua_so():
    ra, _ = chuan_hoa(_docx(["Điều 1. A", "Điều 3. C"]))
    assert [p.text for p in Document(io.BytesIO(ra)).paragraphs] == ["Điều 1. A", "Điều 3. C"]
