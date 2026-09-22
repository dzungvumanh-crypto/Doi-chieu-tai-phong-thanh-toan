"""Test module Chấm TK 459901-1000-000000000 — Cân ITT / Điện KO offline / GD khác.

Quy tắc lấy từ yêu cầu của phòng Thanh toán và bản chấm tay tháng 7/2026 (940 dòng cân ITT,
42 dòng Điện offline, 6 dòng GD khác). Dữ liệu ở đây là DỮ LIỆU TỰ DỰNG bắt chước các mẫu
đó — không chép dữ liệu thật vào repo.

Chạy: .venv\\Scripts\\python.exe -m pytest tests/test_cham459901_000000000.py -v
"""

import io
import logging
import os
import time

import pandas as pd
import pyzipper
import pytest

from backend.services import cham459901_000000000_service as svc

_MK = "matkhau-test-459901-000000000"
_ZIP_MIME = "application/zip"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CHUOI = "Remitting Amount:VND"


@pytest.fixture(autouse=True)
def _mat_khau_zip(monkeypatch):
    monkeypatch.setenv("DOI_CHIEU_ZIP_PASSWORD", _MK)


@pytest.fixture(autouse=True)
def _don_tien_do():
    """Xoá sổ tiến độ của module này trước/sau mỗi test — thay cho việc sửa `tests/conftest.py` (của người khác).
    Một lượt chưa `done` còn sót lại sẽ giữ cửa chốt `phien_doi_chieu` và làm test sau ăn 409."""
    svc._progress.clear()
    yield
    svc._progress.clear()


# ── Dựng dữ liệu ─────────────────────────────────────────────────────────────

def _row(ref, dr=0, cr=0, remark="", **kw):
    d = {
        "TRDATE": "20260701", "TRBRCD": "1000", "USERID": "1000KO", "JOURSEQ": "1",
        "DYTRSEQ": "2", "LOCAC": "459901", "CCY": "VND", "BUSCD": "FX", "UNIT": "IR",
        "TRCD": "", "CUSTOMER": "1000-000000000", "TRTP": "Normal", "REFERENCE": ref,
        "REMARK": remark, "DRAMOUNT": dr, "CRAMOUNT": cr,
    }
    d.update(kw)
    return d


def _df(rows) -> pd.DataFrame:
    """Bảng đã qua bước đọc+lọc (đúng dạng `_phan_loai` nhận vào)."""
    return svc._chuan_hoa(pd.DataFrame(rows))


def _phan(rows):
    can_itt, ko, khac = svc._phan_loai(_df(rows))
    return can_itt, ko, khac


def _refs(d):
    return sorted(d["REFERENCE"])


def _make_zip(rows: list[dict]) -> bytes:
    cols = [*svc._COT_DU_LIEU]
    df = pd.DataFrame([{"CRTDTM": "", **r} for r in rows])[cols]
    buf = io.BytesIO()
    with pyzipper.AESZipFile(buf, "w", compression=pyzipper.ZIP_DEFLATED,
                             encryption=pyzipper.WZ_AES) as zf:
        zf.setpassword(_MK.encode())
        zf.writestr("data.csv", df.to_csv(index=False).encode("utf-8-sig"))
    return buf.getvalue()


def _xlsx(path, rows, extra_cols=()):
    df = pd.DataFrame(rows)
    for c in extra_cols:
        df[c] = ""
    df.to_excel(path, index=False)
    return path


# ── Tên file tồn ─────────────────────────────────────────────────────────────

class TestNhanDienFileTon:
    @pytest.mark.parametrize("ten", [
        "459-mã 0.xlsx",          # tên thật người dùng đang dùng
        "459_mã 0.xlsx",          # tên ghi trong yêu cầu
        "459_TON_T7.xlsx", "459_ton_thang8.xlsx", "459 ma 0.xlsx", "459_MÃ_0_T7.xlsx",
        "459_TON.xls",
    ])
    def test_la_file_ton(self, ten):
        assert svc.classify_upload_filename(ten) == "ton"

    @pytest.mark.parametrize("ten", [
        "~$459_TON.xlsx", "~$459-mã 0.xlsx",   # file khoá tạm của Office, không phải file tồn
        "GL02_20260731_1000.zip",         # GL02 chính — không qua hàm này
        "GL02_20260731_1000.xlsx",
        "1000_gl02_20260701_459901.xlsx", # có 459 nhưng không có ton / mã 0
        # bỏ dấu rồi "tổng"→"tong", "Boston", "Anton" đều chứa chuỗi "ton" nhưng KHÔNG phải từ "ton" (PR review)
        "GL02 459901 Tổng hợp T8.xlsx", "459901_tong_hop.xlsx", "459901_Boston.xlsx", "459901 - Anton.xlsx",
        "459-mã 05.xlsx",                 # "mã 05" không phải "mã 0"
        "459_TON.zip",                    # tồn chỉ nhận Excel
        "ma 0.xlsx",                      # thiếu 459
        "readme.txt",
    ])
    def test_khong_phai_file_ton(self, ten):
        assert svc.classify_upload_filename(ten) is None


# ── Cân ITT ──────────────────────────────────────────────────────────────────

class TestCanITT:
    def test_cung_reference_no_bang_co_thi_can(self):
        can, ko, khac = _phan([
            _row("ITT1", cr=900_000, remark=f"{_CHUOI}900000"),
            _row("ITT1", dr=900_000, remark="AGRIBANK NINH THUAN", USERID="HQNTHANH"),
        ])
        assert _refs(can) == ["ITT1", "ITT1"] and ko.empty and khac.empty

    def test_cung_reference_lech_tien_khong_can(self):
        can, _, khac = _phan([
            _row("ITT1", cr=900_000), _row("ITT1", dr=899_999),
        ])
        assert can.empty and len(khac) == 2

    def test_dong_le_khong_phai_mot_cap(self):
        can, _, khac = _phan([_row("ITT1", dr=0, cr=0)])
        assert can.empty and len(khac) == 1

    def test_nhom_bon_dong_can_thi_lay_ca_bon(self):
        # bản chấm tay tháng 7 có đúng 1 nhóm 4 dòng
        can, _, _ = _phan([
            _row("ITT1", dr=100), _row("ITT1", dr=50),
            _row("ITT1", cr=120), _row("ITT1", cr=30),
        ])
        assert len(can) == 4

    def test_reference_khac_nhau_khong_gom_chung(self):
        # Nợ 500 ở REF A và Có 500 ở REF B: KHÔNG phải cân ITT (đó là việc của KO/khác)
        can, _, _ = _phan([_row("A", dr=500), _row("B", cr=500)])
        assert can.empty

    def test_reference_rong_khong_bi_gom_thanh_mot_nhom(self):
        # REFERENCE rỗng không được coi là "cùng REFERENCE" → không thành Cân ITT. Hai dòng
        # vẫn ghép nhau theo số tiền ở bước KO, nên bị đánh dấu "cần soát lại".
        can, ko, khac = _phan([_row("", dr=500), _row("", cr=500)])
        assert can.empty and khac.empty and len(ko) == 2
        assert (ko["GHI_CHU"] == svc.GHI_CHU_KO_THIEU_CHUOI).all()

    def test_reference_rong_le_ve_gd_khac(self):
        can, ko, khac = _phan([_row("", dr=500), _row("", cr=700)])
        assert can.empty and ko.empty and len(khac) == 2

    def test_reference_co_khoang_trang_van_cung_nhom(self):
        can, _, _ = _phan([_row("ITT1 ", dr=10), _row(" ITT1", cr=10)])
        assert len(can) == 2

    def test_cap_huy_cung_reference_cung_phia_dau_nguoc_vao_can_itt(self):
        """Yêu cầu chỉ nói "cùng REFERENCE, Tổng DRAMOUNT = Tổng CRAMOUNT" — không đòi đủ hai phía. Cancel (−X) và
        Normal (+X) cùng REFERENCE, cùng một phía: tổng triệt tiêu → Cân ITT, không làm bẩn file Điện KO."""
        can, ko, khac = _phan([
            _row("ITT1", cr=500), _row("ITT1", cr=-500, TRTP="Cancel"),
        ])
        assert len(can) == 2 and ko.empty and khac.empty

    def test_cung_phia_dau_nguoc_o_no_cung_can(self):
        can, _, _ = _phan([_row("ITT1", dr=500), _row("ITT1", dr=-500)])
        assert len(can) == 2

    def test_nhom_toan_dong_khong_khong_phai_mot_cap(self):
        can, _, khac = _phan([_row("ITT1", dr=0, cr=0), _row("ITT1", dr=0, cr=0)])
        assert can.empty and len(khac) == 2

    def test_mot_dong_co_ca_no_va_co_bang_nhau_van_la_dong_le(self):
        # n >= 2: một dòng có Nợ = Có > 0 không phải "một cặp giao dịch" (cả Cân ITT lẫn KO)
        can, ko, khac = _phan([_row("Z", dr=5, cr=5)])
        assert can.empty and ko.empty and len(khac) == 1

    def test_tong_lon_khong_lech_do_so_thuc(self):
        # 7.465.869.595.881 — cỡ số thật của TK, float vẫn phải cộng đúng
        can, _, _ = _phan([_row("N1", dr=7_465_869_595_881), _row("N1", cr=7_465_869_595_881)])
        assert len(can) == 2


# ── Điện KO offline ──────────────────────────────────────────────────────────

class TestDienKO:
    def test_cr_co_chuoi_ghep_dr_cung_so_tien_khac_reference(self):
        """Vế Có KO (REFERENCE ...ITT) ghép vế Nợ (REFERENCE ...OTT) theo SỐ TIỀN."""
        can, ko, khac = _phan([
            _row("1000ITT1", cr=3_164_119_200, remark=f"{_CHUOI}3164119200", USERID="1000KC"),
            _row("1000OTT1", dr=3_164_119_200, remark="Trung tâm Thẻ Agribank", USERID="HQHAPT"),
        ])
        assert can.empty and khac.empty
        assert _refs(ko) == ["1000ITT1", "1000OTT1"]
        assert (ko["GHI_CHU"] == "").all()           # có chuỗi → không cần soát lại

    def test_nhieu_dien_cung_so_tien_can_hai_ben(self):
        # 12.120.000 xuất hiện 2 lần ở mỗi phía (ngày 14 và 31)
        _, ko, khac = _phan([
            _row("I1", cr=12_120_000, remark=f"{_CHUOI}12120000"),
            _row("I2", cr=12_120_000, remark=f"{_CHUOI}12120000"),
            _row("O1", dr=12_120_000), _row("O2", dr=12_120_000),
        ])
        assert len(ko) == 4 and khac.empty

    def test_nhom_khong_can_ca_nhom_ve_gd_khac_va_ghi_chu_nghi_ngo(self):
        # 2 điện Có nhưng chỉ 1 Nợ đối ứng → Tổng Nợ ≠ Tổng Có → không tự chốt
        _, ko, khac = _phan([
            _row("I1", cr=100_000, remark=f"{_CHUOI}100000"),
            _row("I2", cr=100_000, remark=f"{_CHUOI}100000"),
            _row("O1", dr=100_000),
        ])
        assert ko.empty and len(khac) == 3
        # chỉ 2 dòng mang chuỗi bị nghi ngờ; dòng Nợ thường không bị gắn ghi chú
        assert (khac["GHI_CHU"] == svc.GHI_CHU_NGHI_KO).sum() == 2
        assert khac[khac["REFERENCE"] == "O1"]["GHI_CHU"].iloc[0] == ""

    def test_dien_ko_le_khong_co_doi_ung_ve_gd_khac(self):
        _, ko, khac = _phan([_row("I1", cr=453_951, remark=f"{_CHUOI}453951")])
        assert ko.empty and len(khac) == 1
        assert khac["GHI_CHU"].iloc[0] == svc.GHI_CHU_NGHI_KO

    def test_cap_dieu_chinh_khong_chuoi_van_vao_ko_nhung_bi_danh_dau(self):
        """Bản chấm tay tháng 7/2026 xếp cặp NAPAS + bút toán điều chỉnh số âm vào Điện offline
        dù KHÔNG có chuỗi — chương trình làm theo mẫu nhưng ghi chú để người chấm soát lại."""
        _, ko, khac = _phan([
            _row("1000ITT5671", dr=7_465_869_595_881, remark="NAPAS", USERID="HQHAPT"),
            _row("1000GEO1589", dr=-7_465_869_595_881, remark="Hạch toán điều chỉnh", USERID="HQHAPT"),
        ])
        assert khac.empty and len(ko) == 2
        assert (ko["GHI_CHU"] == svc.GHI_CHU_KO_THIEU_CHUOI).all()

    def test_chuoi_khong_phan_biet_hoa_thuong(self):
        _, ko, _ = _phan([
            _row("I1", cr=500, remark="remitting amount:vnd500"),
            _row("O1", dr=500),
        ])
        assert len(ko) == 2 and (ko["GHI_CHU"] == "").all()

    def test_khong_co_dung_sai_o_khoa_ko_so_tien_mot_dong(self):
        # số tiền 1 đồng: 2 Nợ + 1 Có → Tổng Nợ 2 ≠ Tổng Có 1. Dung sai "<= 1" sẽ coi là cân (đột biến M8 của phản biện)
        _, ko, khac = _phan([_row("A", dr=1), _row("B", dr=1), _row("C", cr=1)])
        assert ko.empty and len(khac) == 3

    def test_dong_khong_tien_khong_bi_gom_thanh_nhom_ko(self):
        _, ko, khac = _phan([_row("Z1"), _row("Z2")])
        assert ko.empty and len(khac) == 2

    def test_dong_da_can_itt_khong_bi_ko_lay_lai(self):
        # ITT1 tự cân (cùng REF) → Cân ITT; 2 dòng KO còn lại vẫn ghép nhau theo số tiền
        can, ko, khac = _phan([
            _row("ITT1", cr=700, remark=f"{_CHUOI}700"), _row("ITT1", dr=700),
            _row("I2", cr=700, remark=f"{_CHUOI}700"), _row("O2", dr=700),
        ])
        assert _refs(can) == ["ITT1", "ITT1"]
        assert _refs(ko) == ["I2", "O2"] and khac.empty

    def test_cung_so_tien_nhung_khong_can_thi_khong_ket_luan_bua(self):
        # Nợ 500 và Nợ 500 (cùng phía): Tổng Nợ 1000 ≠ Tổng Có 0
        _, ko, khac = _phan([_row("A", dr=500), _row("B", dr=500)])
        assert ko.empty and len(khac) == 2


# ── Tổng thể ─────────────────────────────────────────────────────────────────

def test_ba_nhom_cong_lai_dung_bang_tong_va_khong_trung_dong():
    rows = [
        _row("ITT1", cr=10), _row("ITT1", dr=10),                      # cân ITT
        _row("I2", cr=99, remark=f"{_CHUOI}99"), _row("O2", dr=99),    # KO
        _row("K1", cr=14_441_329),                                     # khác
        _row("K2", dr=989_304_217_800, remark="NAPAS"),                # khác
    ]
    can, ko, khac = _phan(rows)
    assert (len(can), len(ko), len(khac)) == (2, 2, 2)
    assert set(_refs(can)) | set(_refs(ko)) | set(_refs(khac)) == {
        "ITT1", "I2", "O2", "K1", "K2",
    }


def test_tong_no_co_moi_nhom_can_tru_gd_khac():
    can, ko, _ = _phan([
        _row("ITT1", cr=10), _row("ITT1", dr=10),
        _row("I2", cr=99, remark=f"{_CHUOI}99"), _row("O2", dr=99),
        _row("K1", cr=5),
    ])
    for d in (can, ko):
        assert d["DRAMOUNT"].sum() == d["CRAMOUNT"].sum()


# ── Đọc file + lọc tài khoản + ghép tồn ──────────────────────────────────────

class TestProcessFiles:
    def _chay(self, tmp_path, monkeypatch, gl02_rows, ton_rows=None, ton_extra=()):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        gl = _xlsx(tmp_path / "GL02_20260731_1000.xlsx", gl02_rows)
        ton = None
        if ton_rows is not None:
            ton = ("459-mã 0.xlsx", _xlsx(tmp_path / "459-ma-0.xlsx", ton_rows, ton_extra))
        return svc.process_files([("GL02_20260731_1000.xlsx", gl)], None, ton)

    def test_loc_bo_dong_cua_so_khac_usd_va_tai_khoan_khac(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [
            _row("ITT1", cr=10), _row("ITT1", dr=10),
            _row("X1", dr=5, CUSTOMER="1000-000007709"),   # sổ của module kia
            _row("X2", dr=5, CCY="USD"),
            _row("X3", dr=5, LOCAC="502001"),
        ])
        assert r["can_itt_rows"] == 2 and r["filtered_rows"] == 3
        assert r["total_rows"] == 5 and r["ko_rows"] == 0 and r["khac_rows"] == 0

    def test_ghep_dong_ton_de_can_voi_dong_thang_nay(self, tmp_path, monkeypatch):
        # Dòng Có tồn từ tháng 6 (file tồn có cột tay `chấm`, không có CRTDTM) cân với Nợ tháng 7
        r = self._chay(
            tmp_path, monkeypatch,
            [_row("ITT6457", dr=169_911_841, TRDATE="20260702")],
            ton_rows=[
                _row("ITT6457", cr=169_911_841, TRDATE="20260629"),
                _row("ITT5624", cr=14_441_329, TRDATE="20260527"),   # tồn chưa xử lý xong
            ],
            ton_extra=("chấm",),
        )
        assert r["ton_rows_added"] == 2
        assert (r["can_itt_rows"], r["ko_rows"], r["khac_rows"]) == (2, 0, 1)
        assert r["total_rows"] == 3

    def test_thieu_file_ton_van_chay_binh_thuong(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("ITT1", cr=10), _row("ITT1", dr=10)])
        assert r["ton_rows_added"] == 0 and r["can_itt_rows"] == 2
        assert r["ton_provided"] is False       # để màn hình cảnh báo "KHÔNG có file tồn"

    def test_thieu_ton_lam_dong_doi_ung_roi_ve_gd_khac(self, tmp_path, monkeypatch):
        """Tháng 8/2026: không chọn 459_TON.xlsx → 426/28/4 thay vì 428/30/6. Dòng Nợ 3.000.000 của
        tháng này mất cặp Có tồn nên rơi về GD khác. Khoá lại đúng cơ chế đó."""
        gl = [_row("ITT7072", dr=3_000_000, TRDATE="20260811")]
        ton = [_row("ITT7072", cr=3_000_000, TRDATE="20260716")]
        khong_ton = self._chay(tmp_path, monkeypatch, gl)
        co_ton = self._chay(tmp_path, monkeypatch, gl, ton_rows=ton)
        assert (khong_ton["can_itt_rows"], khong_ton["khac_rows"]) == (0, 1)
        assert (co_ton["can_itt_rows"], co_ton["khac_rows"]) == (2, 0)
        assert co_ton["ton_provided"] is True

    def test_file_ton_khong_co_dong_cua_tai_khoan_khac_voi_khong_chon(self, tmp_path, monkeypatch):
        r = self._chay(
            tmp_path, monkeypatch, [_row("ITT1", cr=10), _row("ITT1", dr=10)],
            ton_rows=[_row("X", cr=5, CUSTOMER="1000-000007709")],   # tồn của sổ khác
        )
        assert r["ton_provided"] is True and r["ton_rows_added"] == 0

    def test_file_tong_hop_ket_qua_ghi_du_ba_file_excel(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [
            _row("ITT1", cr=10), _row("ITT1", dr=10), _row("K1", cr=5),
        ])
        thu_muc = tmp_path / "out" / r["token"]
        assert sorted(p.name for p in thu_muc.glob("*.xlsx")) == ["can_itt.xlsx", "khac.xlsx", "ko.xlsx"]
        # đọc lại đúng số dòng (+1 dòng TỔNG CỘNG do _write_excel thêm)
        d = pd.read_excel(thu_muc / "can_itt.xlsx", skiprows=1, engine="calamine")
        assert len(d) == 2 + 1

    def test_ghi_chu_ghi_ra_file_khac(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("I1", cr=453_951, remark=f"{_CHUOI}453951")])
        assert r["khac_nghi_ko_rows"] == 1
        d = pd.read_excel(tmp_path / "out" / r["token"] / "khac.xlsx", skiprows=1,
                          engine="calamine", dtype=str, keep_default_na=False)
        assert svc.GHI_CHU_NGHI_KO in set(d["GHI_CHU"])

    def test_khong_co_dong_nao_cua_tai_khoan_bao_loi_de_hieu(self, tmp_path, monkeypatch):
        with pytest.raises(svc.InputError, match="1000-000000000"):
            self._chay(tmp_path, monkeypatch, [_row("X", dr=5, CUSTOMER="1000-000007709")])

    def test_khong_chon_file_nao(self):
        with pytest.raises(svc.InputError):
            svc.process_files([])

    def test_zip_ma_hoa_va_excel_tron_trong_mot_luot(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        z = tmp_path / "GL02_thang7.zip"
        z.write_bytes(_make_zip([_row("ITT1", cr="10"), _row("ITT1", dr="10")]))
        x = _xlsx(tmp_path / "them.xlsx", [_row("I2", cr=99, remark=f"{_CHUOI}99"), _row("O2", dr=99)])
        r = svc.process_files([("GL02_thang7.zip", z), ("them.xlsx", x)])
        assert (r["can_itt_rows"], r["ko_rows"], r["khac_rows"]) == (2, 2, 0)

    def test_file_hong_bao_loi_kem_ten_file(self, tmp_path):
        hong = tmp_path / "GL02.zip"
        hong.write_bytes(b"khong phai zip")
        with pytest.raises(svc.InputError, match="GL02.zip"):
            svc.process_files([("GL02.zip", hong)])

    def test_chan_dung_thoi_diem_bam_dung(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        token = svc.init_progress()
        assert svc.cancel_progress(token)
        gl = _xlsx(tmp_path / "GL02.xlsx", [_row("ITT1", cr=10), _row("ITT1", dr=10)])
        with pytest.raises(svc._Cancelled):
            svc.process_files([("GL02.xlsx", gl)], token)
        assert not (tmp_path / "out").exists() or not any((tmp_path / "out").glob("*/can_itt.xlsx"))


# ── Số tiền so sánh CHÍNH XÁC (Decimal), không round()/== trên float ──────────────

class TestSoTienChinhXac:
    """Checklist mục E (PR #60, #61). Mỗi ca ở đây là ca mà bản `(sum_dr - sum_cr).round(2) == 0` trên float
    cho kết quả SAI — số tiền VND thật thì nguyên nên không lộ, nhưng phép so sánh không được phép có dung sai."""

    def test_lech_bon_phan_van_khong_duoc_coi_la_can(self):
        # float: |10.0004 - 10.0000| = 0.0004 → round(2) = 0.0 → "cân". Decimal: lệch thật.
        can, _, khac = _phan([_row("A", dr="10.0004"), _row("A", cr="10.0000")])
        assert can.empty and len(khac) == 2

    def test_0_1_cong_0_2_bang_dung_0_3(self):
        # float: 0.1 + 0.2 = 0.30000000000000004 ≠ 0.3. Decimal dựng từ chuỗi: đúng bằng nhau.
        can, _, _ = _phan([_row("A", dr="0.1"), _row("A", dr="0.2"), _row("A", cr="0.3")])
        assert len(can) == 3

    def test_so_lon_hon_2_mu_53_khong_bi_lam_tron(self):
        # 2^53 + 1 = 9007199254740993: float làm tròn về ...992 nên hai số này "bằng nhau" trong float
        can, _, khac = _phan([_row("A", dr="9007199254740993"), _row("A", cr="9007199254740992")])
        assert can.empty and len(khac) == 2

    def test_khoa_nhom_ko_khong_lam_tron_so_tien(self):
        # bản cũ gom theo round(0): 5000.0004 và 5000 vào cùng nhóm rồi "cân" nhờ dung sai
        _, ko, khac = _phan([
            _row("I", cr="5000.0004", remark=f"{_CHUOI}5000"), _row("O", dr="5000"),
        ])
        assert ko.empty and len(khac) == 2

    def test_so_am_va_so_thuc_kieu_float_van_dung(self):
        can, _, _ = _phan([_row("N", dr=-7.5), _row("N", cr=-7.5)])          # đầu vào kiểu số, không phải chuỗi
        assert len(can) == 2

    def test_cot_phu_decimal_khong_lot_ra_file_excel(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        gl = _xlsx(tmp_path / "GL02.xlsx", [_row("A", cr=10), _row("A", dr=10)])
        r = svc.process_files([("GL02.xlsx", gl)])
        d = pd.read_excel(tmp_path / "out" / r["token"] / "can_itt.xlsx", skiprows=1, engine="calamine")
        assert list(d.columns) == ["STT", *svc.OUTPUT_COLS]
        assert not any(str(c).startswith("_") for c in d.columns)
        assert d["DRAMOUNT"].iloc[:-1].sum() == 10                           # vẫn là số, cộng đúng


# ── Những chỗ từng CÂM LẶNG: sai/thiếu dữ liệu vào mà không báo gì ─────────────

class TestKhongCamLang:
    def _chay(self, tmp_path, monkeypatch, gl_rows, ton_rows=None, cot_bo=(), ten="GL02.xlsx"):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        p = tmp_path / ten
        df = pd.DataFrame(gl_rows)
        df.drop(columns=list(cot_bo)).to_excel(p, index=False)
        ton = None
        if ton_rows is not None:
            ton = ("459_TON.xlsx", _xlsx(tmp_path / "ton.xlsx", ton_rows))
        return svc.process_files([(ten, p)], None, ton)

    # 1) số tiền dạng chữ
    def test_so_tien_dang_chu_co_dau_phay_bao_loi_khong_doi_thanh_0(self, tmp_path, monkeypatch):
        """Trước đây "1,000" → 0 lặng lẽ: cặp Nợ/Có bằng nhau không còn cân, rơi hết xuống GD khác."""
        with pytest.raises(svc.InputError, match=r"không phải số.*1,000") as e:
            self._chay(tmp_path, monkeypatch, [_row("ITT9", cr="1,000"), _row("ITT9", dr="1,000")])
        msg = str(e.value)
        assert "GL02.xlsx" in msg and "ITT9" in msg
        assert "DRAMOUNT" in msg and "CRAMOUNT" in msg          # báo CẢ HAI cột một lần, khỏi sửa xong mới thấy cột kia

    def test_chi_mot_cot_sai_chi_nhac_cot_do(self, tmp_path, monkeypatch):
        with pytest.raises(svc.InputError) as e:
            self._chay(tmp_path, monkeypatch, [_row("ITT9", cr="abc"), _row("ITT9", dr=5)])
        assert "CRAMOUNT" in str(e.value) and "DRAMOUNT" not in str(e.value)

    def test_o_tien_de_trong_van_la_0_khong_bao_loi(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [
            _row("ITT1", cr=10), _row("ITT1", dr=10),
            _row("Z1", dr=None, cr=None),          # ô trống — bình thường, không phải lỗi
        ])
        assert r["can_itt_rows"] == 2 and r["khac_rows"] == 1 and r["canh_bao"] == []

    def test_so_am_va_so_thap_phan_van_doc_dung(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("N1", dr=-7.5), _row("N1", cr=-7.5)])
        assert r["can_itt_rows"] == 2

    def test_o_tien_dang_x_yyy_mo_ho_bao_loi_khong_doc_thanh_1(self, tmp_path, monkeypatch):
        """"1.000" là cách viết Việt Nam của "một nghìn" nhưng máy đọc thành 1 (sai gấp 1000 lần, không báo) —
        đúng bẫy checklist mục E. Phản biện tái hiện được; nay chặn vì mơ hồ."""
        with pytest.raises(svc.InputError, match=r"mơ hồ.*1\.000") as e:
            self._chay(tmp_path, monkeypatch, [_row("A", cr="1.000"), _row("A", dr="1.000")])
        assert "gấp 1000 lần" in str(e.value)

    def test_so_thap_phan_binh_thuong_van_doc_dung(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("A", cr="1000.00"), _row("A", dr="1000.00"),
                                                _row("B", cr="7.25"), _row("B", dr="7.25")])
        assert r["can_itt_rows"] == 4

    # 2) thiếu cột REFERENCE
    def test_thieu_cot_reference_bao_loi_thay_vi_khong_can_cap_nao(self, tmp_path, monkeypatch):
        with pytest.raises(svc.InputError, match="thiếu cột REFERENCE") as e:
            self._chay(tmp_path, monkeypatch, [_row("A", cr=10), _row("A", dr=10)], cot_bo=("REFERENCE",))
        assert "GL02.xlsx" in str(e.value)

    # 3) trùng dòng
    def test_cung_du_lieu_vao_hai_lan_duoc_canh_bao_va_khong_tu_xoa(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        rows = [_row("A", cr=10), _row("A", dr=10), _row("B", cr=5)]
        a = _xlsx(tmp_path / "a.xlsx", rows)
        b = _xlsx(tmp_path / "b.xlsx", rows)                  # y hệt, tên khác → API không chặn được
        r = svc.process_files([("a.xlsx", a), ("b.xlsx", b)])
        assert len(r["canh_bao"]) == 1 and "3 dòng trùng hoàn toàn" in r["canh_bao"][0]
        assert (r["can_itt_rows"], r["khac_rows"]) == (4, 2)  # KHÔNG tự xoá — người dùng quyết

    def test_file_ton_trung_voi_gl02_cung_bi_canh_bao(self, tmp_path, monkeypatch):
        dong = _row("K1", cr=5, TRDATE="20260701")
        r = self._chay(tmp_path, monkeypatch, [dong], ton_rows=[dong])
        assert any("trùng hoàn toàn" in c for c in r["canh_bao"])

    def test_du_lieu_sach_khong_co_canh_bao_nao(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [
            _row("ITT1", cr=10, JOURSEQ="1"), _row("ITT1", dr=10, JOURSEQ="2"),
            _row("K1", cr=5, JOURSEQ="3"),
        ], ton_rows=[_row("K0", cr=7, TRDATE="20260630")])
        assert r["canh_bao"] == []                             # không báo nhầm

    def test_hai_but_toan_khac_jourseq_khong_bi_coi_la_trung(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("A", cr=10, JOURSEQ="1"), _row("A", cr=10, JOURSEQ="2"),
                                                _row("A", dr=20, JOURSEQ="3")])
        assert r["canh_bao"] == [] and r["can_itt_rows"] == 3

    # 4) Excel bị cắt ở trần dòng
    def test_file_excel_sat_tran_dong_bi_canh_bao_cat_bot(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "_TRAN_DONG_EXCEL", 4)        # ngưỡng nhỏ để thử rẻ
        r = self._chay(tmp_path, monkeypatch, [_row(f"R{i}", dr=i, JOURSEQ=str(i)) for i in range(1, 6)])
        assert any("1.048.576" in c and "cắt bớt" in c and "GL02.xlsx" in c for c in r["canh_bao"])

    def test_zip_dai_bang_tran_excel_khong_bi_bao_nham(self, tmp_path, monkeypatch):
        """Zip/CSV không bị trần Excel — một CSV gốc GL02 dài ~1 triệu dòng là chuyện thường."""
        monkeypatch.setattr(svc, "_TRAN_DONG_EXCEL", 4)
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        z = tmp_path / "GL02.zip"
        z.write_bytes(_make_zip([_row(f"R{i}", dr=str(i), JOURSEQ=str(i)) for i in range(1, 6)]))
        assert svc.process_files([("GL02.zip", z)])["canh_bao"] == []

    # 5) GL02 không có dòng nào của TK nhưng có tồn
    def test_gl02_khong_co_dong_nao_cua_tk_ma_co_ton_thi_canh_bao(self, tmp_path, monkeypatch):
        r = self._chay(
            tmp_path, monkeypatch, [_row("Z", dr=9, CUSTOMER="1000-000007709")],
            ton_rows=[_row("K1", cr=5, TRDATE="20260630")],
        )
        assert r["ton_rows_added"] == 1 and r["khac_rows"] == 1
        assert any("không có dòng nào của TK" in c and "1 dòng tồn" in c for c in r["canh_bao"])

    def test_canh_bao_nam_trong_ket_qua_tra_ve_giao_dien(self, tmp_path, monkeypatch):
        r = self._chay(tmp_path, monkeypatch, [_row("A", cr=10), _row("A", dr=10)])
        assert isinstance(r["canh_bao"], list)                 # UI lặp qua r["canh_bao"], thiếu khoá là im lặng


# ── Cột STT ở các file xuất ──────────────────────────────────────────────────

class TestCotSTT:
    def _chay(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path / "out")
        gl = _xlsx(tmp_path / "GL02.xlsx", [
            _row("ITT1", cr=10), _row("ITT1", dr=10),                     # cân ITT (2)
            _row("ITT2", cr=20), _row("ITT2", dr=20),
            _row("I3", cr=99, remark=f"{_CHUOI}99"), _row("O3", dr=99),   # KO (2)
            _row("K1", cr=5), _row("K2", cr=6), _row("K3", cr=7),         # khác (3)
        ])
        r = svc.process_files([("GL02.xlsx", gl)])
        return r, tmp_path / "out" / r["token"]

    @pytest.mark.parametrize("ten_file, so_dong", [("can_itt.xlsx", 4), ("ko.xlsx", 2), ("khac.xlsx", 3)])
    def test_ca_ba_file_co_cot_stt_la_so_lien_tuc(self, tmp_path, monkeypatch, ten_file, so_dong):
        _, thu_muc = self._chay(tmp_path, monkeypatch)
        d = pd.read_excel(thu_muc / ten_file, skiprows=1, engine="calamine")
        assert list(d.columns[:2]) == ["STT", "TRDATE"]          # STT là cột A, dữ liệu gốc lùi sang B…
        assert d.columns[-1] == "GHI_CHU"
        stt = d["STT"].dropna()
        assert stt.tolist() == list(range(1, so_dong + 1))       # 1, 2, 3… là SỐ, không phải chữ
        assert pd.api.types.is_numeric_dtype(d["STT"])
        # dòng cuối là TỔNG CỘNG: STT trống, nhãn nằm ở cột TRDATE (cột A hẹp không đủ chỗ)
        assert pd.isna(d["STT"].iloc[-1]) and d["TRDATE"].iloc[-1] == "TỔNG CỘNG"
        assert len(d) == so_dong + 1

    def test_tong_cong_khong_bi_lech_khi_them_cot(self, tmp_path, monkeypatch):
        _, thu_muc = self._chay(tmp_path, monkeypatch)
        d = pd.read_excel(thu_muc / "can_itt.xlsx", skiprows=1, engine="calamine")
        du_lieu, tong = d.iloc[:-1], d.iloc[-1]
        assert tong["DRAMOUNT"] == du_lieu["DRAMOUNT"].sum() == 30
        assert tong["CRAMOUNT"] == du_lieu["CRAMOUNT"].sum() == 30

    def test_file_mo_duoc_bang_openpyxl_va_bo_loc_bao_het_cot(self, tmp_path, monkeypatch):
        import openpyxl
        _, thu_muc = self._chay(tmp_path, monkeypatch)
        ws = openpyxl.load_workbook(thu_muc / "can_itt.xlsx").active
        assert ws["A2"].value == "STT" and ws["B2"].value == "TRDATE"
        assert ws["A3"].value == 1 and ws["A6"].value == 4       # kiểu số, không phải chuỗi
        assert ws.auto_filter.ref == f"A2:S{4 + 2}"              # 19 cột = A..S, bộ lọc phủ cả cột STT
        assert "A1:S1" in [str(m) for m in ws.merged_cells.ranges]

    def test_stt_lien_tuc_xuyen_cac_sheet_khi_nhom_bi_tach(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "_MAX_DATA_ROWS", 5)            # trần nhỏ để thử tách sheet rẻ
        d = _df([_row(f"R{i}", dr=i) for i in range(1, 13)])
        d["GHI_CHU"] = ""
        out = tmp_path / "tach.xlsx"
        svc._ghi_excel(d, out, "GD khác", "E67E22")
        sheets = pd.read_excel(out, sheet_name=None, skiprows=1, engine="calamine")
        assert len(sheets) == 3
        stt = [s["STT"].dropna().astype(int).tolist() for s in sheets.values()]
        assert stt == [[1, 2, 3, 4, 5], [6, 7, 8, 9, 10], [11, 12]]      # không đánh lại từ 1 ở sheet sau
        # mỗi sheet có dòng tổng của CHÍNH nó (nhãn ở cột B, ghi rõ phần)
        tong = [s["TRDATE"].iloc[-1] for s in sheets.values()]
        assert tong == ["TỔNG CỘNG PHẦN 1/3", "TỔNG CỘNG PHẦN 2/3", "TỔNG CỘNG PHẦN 3/3"]

    def test_ghi_excel_khong_dung_toi_ham_ghi_cua_module_cu(self):
        """Cột STT nằm ở hàm ghi RIÊNG của module này: `cham459901_service._write_excel` (phần của người khác) không
        có tham số nào thêm và vẫn nhận đúng chữ ký gốc."""
        import inspect
        from backend.services import cham459901_service as goc
        assert list(inspect.signature(goc._write_excel).parameters) == ["df", "path", "sheet_name", "hex_color"]
        assert "STT" not in goc.COL_WIDTHS and "STT" not in goc.OUTPUT_COLS

    def test_file_gd_khac_thang_nay_lam_file_ton_thang_sau_van_doc_duoc(self, tmp_path, monkeypatch):
        """Thực tế người chấm dùng GD khác của tháng này làm file TỒN tháng sau (tháng 8: 459_TON.xlsx = 6 dòng
        GD khác tháng 7). File giờ có cột STT + dòng TỔNG CỘNG — vẫn phải đọc và ghép đúng, không sinh dòng ma."""
        r1, thu_muc = self._chay(tmp_path, monkeypatch)
        khac = tmp_path / "459_TON_tu_khac.xlsx"
        khac.write_bytes((thu_muc / "khac.xlsx").read_bytes())
        # tháng sau: dòng Nợ 5 (đối ứng của K1 Có 5 tồn) + không liên quan gì khác
        gl = _xlsx(tmp_path / "GL02_thang_sau.xlsx", [_row("K1", dr=5)])
        r2 = svc.process_files([("GL02_thang_sau.xlsx", gl)], None, ("459_TON_tu_khac.xlsx", khac))
        assert r2["ton_rows_added"] == r1["khac_rows"] == 3       # 3 dòng thật, KHÔNG tính dòng TỔNG CỘNG
        assert r2["can_itt_rows"] == 2                            # K1 tồn (Có 5) cân với K1 tháng sau (Nợ 5)
        assert r2["khac_rows"] == 2 and r2["total_rows"] == 4


# ── Vòng đời job ─────────────────────────────────────────────────────────────

class TestJobLifecycle:
    def test_cancel_va_get_progress(self):
        token = svc.init_progress()
        p = svc.get_progress(token)
        assert p["done"] is False and "cancel_event" not in p
        assert svc.cancel_progress(token) is True
        svc.bo_luot(token)
        assert svc.get_progress(token) is None
        assert svc.cancel_progress(token) is False

    def test_luot_dang_chay_va_het_han(self):
        assert svc.luot_dang_chay() is None
        token = svc.init_progress()
        assert svc.luot_dang_chay()["job_id"] == token
        svc._progress[token]["_ts"] -= svc._TTL_DANG_CHAY + 1   # bỏ dở quá lâu → coi như chết
        assert svc.luot_dang_chay() is None

    def test_don_dep_xoa_ca_tien_do_cu_va_thu_muc_cu(self, tmp_path, monkeypatch):
        """`_cleanup_old_results` phải dọn cả mục `_progress` cũ chứ không chỉ thư mục (đột biến M30)."""
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        cu, moi = svc.init_progress(), svc.init_progress()
        svc._progress[cu]["_ts"] -= 10_000
        thu_muc_cu = tmp_path / "ket-qua-cu"
        thu_muc_cu.mkdir()
        cu_mtime = time.time() - 9_999
        os.utime(thu_muc_cu, (cu_mtime, cu_mtime))
        svc._cleanup_old_results(cutoff=time.time() - 5_000)
        assert cu not in svc._progress and moi in svc._progress
        assert not thu_muc_cu.exists()

    def test_delete_result(self, tmp_path, monkeypatch):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        (tmp_path / "abc").mkdir()
        assert svc.delete_result("abc") is True
        assert svc.delete_result("abc") is False


# ── API ──────────────────────────────────────────────────────────────────────

_API = "/api/cham459901_000000000"


def _wait_done(client, task_token, timeout_s=10):
    deadline = time.time() + timeout_s
    prog = None
    while time.time() < deadline:
        r = client.get(f"{_API}/progress/{task_token}")
        assert r.status_code == 200
        prog = r.json()
        if prog["done"]:
            return prog
        time.sleep(0.05)
    raise AssertionError(f"Job không hoàn thành sau {timeout_s}s: {prog}")


def _cap_itt():
    return [_row("REF1", cr="10"), _row("REF1", dr="10")]


class TestApi:
    def test_xu_ly_zip_va_ton_roi_tai_ba_file(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        ton = _xlsx(tmp_path / "t.xlsx", [_row("REF0", cr=7, TRDATE="20260629")]).read_bytes()

        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_20260731_1000.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
            ("files", ("459-mã 0.xlsx", ton, _XLSX_MIME)),
        ])
        assert r.status_code == 200
        body = r.json()
        assert body["unrecognized"] == [] and body["duplicates"] == {}

        prog = _wait_done(admin_client, body["task_token"])
        assert prog["error"] is None and prog["cancelled"] is False
        res = prog["result"]
        assert (res["can_itt_rows"], res["khac_rows"], res["ton_rows_added"]) == (2, 1, 1)

        for ft in ("can_itt", "ko", "khac"):
            d = admin_client.get(f"{_API}/download/{res['token']}/{ft}")
            assert d.status_code == 200
            # Tên tải về mang mã TK để người dùng phân biệt với module 1000-000007709
            assert "459901-1000-000000000" in d.headers["content-disposition"]

    def test_thieu_gl02_400(self, admin_client):
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("bao_cao.pdf", b"%PDF-1.4", "application/pdf")),
        ])
        assert r.status_code == 400 and "GL02" in r.json()["detail"]

    def test_chi_co_file_ton_khong_co_gl02_400(self, admin_client, tmp_path):
        ton = _xlsx(tmp_path / "t.xlsx", [_row("REF0", cr=7)]).read_bytes()
        r = admin_client.post(f"{_API}/process", files=[("files", ("459-mã 0.xlsx", ton, _XLSX_MIME))])
        assert r.status_code == 400

    def test_file_la_bi_bo_qua_khong_chan_ca_luot(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_1000.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
            ("files", ("ghi_chu.docx", b"x", "application/msword")),
        ])
        assert r.status_code == 200 and r.json()["unrecognized"] == ["ghi_chu.docx"]
        # Chờ job xong TRƯỚC khi test kết thúc — luồng nền chạy tràn sang sau monkeypatch
        # sẽ ghi vào data/temp_* thật
        _wait_done(admin_client, r.json()["task_token"])

    def test_hai_file_ton_canh_bao_khong_chan(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        ton = _xlsx(tmp_path / "t.xlsx", [_row("REF0", cr=7)]).read_bytes()
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_1000.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
            ("files", ("459_TON_T6.xlsx", ton, _XLSX_MIME)),
            ("files", ("459-mã 0.xlsx", ton, _XLSX_MIME)),
        ])
        assert r.status_code == 200
        assert set(r.json()["duplicates"]["ton"]) == {"459_TON_T6.xlsx", "459-mã 0.xlsx"}
        _wait_done(admin_client, r.json()["task_token"])

    def test_file_khoa_office_bi_bo_qua_khong_de_mat_file_ton_that(self, admin_client, monkeypatch, tmp_path):
        """Thư mục tháng 8 của phòng có `~$1000_gl02_..._1.xlsx` (Excel đang mở file). Kéo-thả cả
        thư mục mà mang theo nó thì: `~$459_TON.xlsx` bị nhận là file TỒN và (file cuối thắng)
        đè mất file tồn thật; `~$...gl02.xlsx` thì làm hỏng cả lượt vì không đọc được như Excel."""
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        ton = _xlsx(tmp_path / "t.xlsx", [_row("REF0", cr=7, TRDATE="20260629")]).read_bytes()
        khoa = b"\x00" * 165        # file khoá thật dài 165 byte, không phải Excel

        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("1000_gl02_20260801.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
            ("files", ("459_TON.xlsx", ton, _XLSX_MIME)),
            ("files", ("~$459_TON.xlsx", khoa, _XLSX_MIME)),
            ("files", ("~$1000_gl02_2026080120260831_1.xlsx", khoa, _XLSX_MIME)),
        ])
        assert r.status_code == 200
        body = r.json()
        assert set(body["unrecognized"]) == {"~$459_TON.xlsx", "~$1000_gl02_2026080120260831_1.xlsx"}
        assert body["duplicates"] == {}      # file tồn thật không bị coi là trùng với file khoá

        prog = _wait_done(admin_client, body["task_token"])
        assert prog["error"] is None, prog
        assert prog["result"]["ton_rows_added"] == 1     # file tồn THẬT vẫn được ghép

    def test_chi_co_file_khoa_khong_co_gl02_400(self, admin_client):
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("~$1000_gl02_2026080120260831.xlsx", b"\x00" * 165, _XLSX_MIME)),
        ])
        assert r.status_code == 400 and "GL02" in r.json()["detail"]

    def test_trung_ten_khac_hoa_thuong_cung_bi_chan(self, admin_client, monkeypatch, tmp_path):
        """Đĩa Windows (NTFS) không phân biệt hoa/thường: `GL02.zip` và `gl02.ZIP` là MỘT file, file sau đè
        file trước rồi bị đọc hai lần. Phản biện chỉ ra chỗ này (bản cũ so tên phân biệt hoa/thường)."""
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        z = _make_zip(_cap_itt())
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02.zip", z, _ZIP_MIME)), ("files", ("gl02.ZIP", z, _ZIP_MIME)),
        ])
        assert r.status_code == 400 and "hai lần" in r.json()["detail"]
        assert not list(tmp_path.glob("upload_*")) and svc.luot_dang_chay() is None

    def test_khong_co_gl02_thi_don_sach_thu_muc_va_nha_cua_chot(self, admin_client, monkeypatch, tmp_path):
        """Bỏ dòng `bo_luot` ở nhánh 'không có GL02' thì cửa chốt 409 kẹt tới 4 giờ và rò thư mục upload —
        không test nào canh (đột biến M23/M24 của phản biện)."""
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        r = admin_client.post(f"{_API}/process", files=[("files", ("bao_cao.pdf", b"%PDF-1.4", "application/pdf"))])
        assert r.status_code == 400
        assert not list(tmp_path.glob("upload_*")), "thư mục upload phải được dọn"
        assert svc._progress == {} and svc.luot_dang_chay() is None

    def test_vuot_tran_dung_luong_bao_413_va_don_sach(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        monkeypatch.setattr("backend.api.cham459901_000000000.MAX_REQUEST_BYTES", 50)   # trần nhỏ để thử rẻ
        r = admin_client.post(f"{_API}/process", files=[("files", ("GL02.zip", b"x" * 500, _ZIP_MIME))])
        assert r.status_code == 413 and "vượt quá" in r.json()["detail"]
        assert not list(tmp_path.glob("upload_*")) and svc.luot_dang_chay() is None

    def test_tao_thu_muc_loi_thi_van_nha_cua_chot(self, admin_client, monkeypatch, tmp_path):
        """`mkdir` lỗi (đĩa đầy/quyền) ở NGOÀI khối try thì mục tiến độ mồ côi giữ cửa chốt 409 tới 4 giờ."""
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)

        def _hong(_token):
            raise OSError("dia day")

        monkeypatch.setattr(svc, "tao_thu_muc_upload", _hong)
        with pytest.raises(OSError):
            admin_client.post(f"{_API}/process", files=[("files", ("GL02.zip", _make_zip(_cap_itt()), _ZIP_MIME))])
        assert svc._progress == {} and svc.luot_dang_chay() is None

    def test_chon_trung_mot_file_hai_lan_400(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        z = _make_zip(_cap_itt())
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_1000.zip", z, _ZIP_MIME)),
            ("files", ("GL02_1000.zip", z, _ZIP_MIME)),
        ])
        assert r.status_code == 400 and "hai lần" in r.json()["detail"]

    def test_chan_luot_thu_hai_khi_dang_chay_409(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        svc.init_progress()                       # một lượt đang chiếm
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_1000.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
        ])
        assert r.status_code == 409

    def test_cancel_progress_delete_download_khong_ton_tai(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        assert admin_client.post(f"{_API}/cancel/khong-ton-tai").status_code == 404
        assert admin_client.get(f"{_API}/progress/khong-ton-tai").status_code == 404
        assert admin_client.delete(f"{_API}/result/khong-ton-tai").status_code == 404
        assert admin_client.get(f"{_API}/download/khong-ton-tai/khac").status_code == 404
        assert admin_client.get(f"{_API}/download/any/khong_hop_le").status_code == 400

    def test_xoa_ket_qua_roi_tai_404(self, admin_client, monkeypatch, tmp_path):
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        r = admin_client.post(f"{_API}/process", files=[
            ("files", ("GL02_1000.zip", _make_zip(_cap_itt()), _ZIP_MIME)),
        ])
        token = _wait_done(admin_client, r.json()["task_token"])["result"]["token"]
        assert admin_client.get(f"{_API}/download/{token}/can_itt").status_code == 200
        assert admin_client.delete(f"{_API}/result/{token}").status_code == 200
        assert admin_client.get(f"{_API}/download/{token}/can_itt").status_code == 404

    def test_token_khong_thoat_duoc_khoi_thu_muc_tam(self, admin_client, monkeypatch, tmp_path):
        temp_dir = tmp_path / "temp"
        temp_dir.mkdir()
        monkeypatch.setattr(svc, "TEMP_DIR", temp_dir)
        bi_mat = tmp_path / "bi_mat"
        bi_mat.mkdir()
        (bi_mat / "khac.xlsx").write_bytes(b"khong phai excel that")
        (bi_mat / "quan_trong.txt").write_text("giu lai")

        assert admin_client.delete(f"{_API}/result/..%2Fbi_mat").status_code == 404
        assert (bi_mat / "quan_trong.txt").exists()
        assert admin_client.get(f"{_API}/download/..%2Fbi_mat/khac").status_code == 404

    def test_token_khong_thoat_duoc_bang_dau_gach_nguoc(self, admin_client, monkeypatch, tmp_path):
        """Checklist mục A: thử cả `\\` chứ không chỉ `/`. Uvicorn chặn `%2F` nhưng `%5C` thì không, và
        `pathlib` trên Windows coi `\\` là dấu phân cách nên `Path('x') / '..\\\\bi_mat'` thoát ra ngoài."""
        temp_dir = tmp_path / "temp"
        temp_dir.mkdir()
        monkeypatch.setattr(svc, "TEMP_DIR", temp_dir)
        bi_mat = tmp_path / "bi_mat"
        bi_mat.mkdir()
        (bi_mat / "khac.xlsx").write_bytes(b"khong phai excel that")
        (bi_mat / "quan_trong.txt").write_text("giu lai")

        for kieu in ("..%5Cbi_mat", "..%5C..%5Cbi_mat", "%5C..%5Cbi_mat", "..%2F..%5Cbi_mat"):
            assert admin_client.delete(f"{_API}/result/{kieu}").status_code in (404, 405), kieu
            assert admin_client.get(f"{_API}/download/{kieu}/khac").status_code in (404, 405), kieu
        assert (bi_mat / "quan_trong.txt").exists() and (bi_mat / "khac.xlsx").exists()

    def test_khong_co_cua_chon_thu_muc_server(self, admin_client, tmp_path):
        r = admin_client.post(f"{_API}/process_folder", json={"folder_path": str(tmp_path)})
        assert r.status_code in (404, 405)


class TestHaiModuleKhongDungChung:
    """Hai mã CUSTOMER là hai sổ khác nhau — không được lẫn thư mục, tiến độ hay mã quyền."""

    def test_thu_muc_tam_va_bang_tien_do_rieng(self):
        from backend.services import cham459901_service as goc
        assert svc.TEMP_DIR != goc.TEMP_DIR
        assert svc._progress is not goc._progress
        assert svc.FILTER_CUSTOMER == "1000-000000000" != goc.FILTER_CUSTOMER

    def test_tien_do_cua_module_nay_khong_hien_o_module_kia(self, admin_client):
        token = svc.init_progress()
        assert admin_client.get(f"{_API}/progress/{token}").status_code == 200
        assert admin_client.get(f"/api/cham459901/progress/{token}").status_code == 404

    def test_hai_ma_quyen_rieng_va_ten_module_moi_mang_ma_tk(self):
        from backend.core.features import FEATURES
        for ma in ("menu.cham_459901_000000000", "cham_459901_000000000.process",
                   "menu.cham_459901", "cham_459901.process"):
            assert ma in FEATURES
        assert "000000000" in FEATURES["menu.cham_459901_000000000"]     # tên module mới có mã TK để phân biệt
        assert FEATURES["menu.cham_459901"] != FEATURES["menu.cham_459901_000000000"]

    def test_nhan_nhat_ky_khong_bi_module_kia_nuot(self):
        from backend.services.audit_labels import describe_work
        # "/api/cham459901" là tiền tố chuỗi của "/api/cham459901_000000000": dòng dài hơn phải đứng trước
        moi = describe_work("POST", "/api/cham459901_000000000/cancel/x")
        cu = describe_work("POST", "/api/cham459901/cancel/x")
        assert "000000000" in moi and "000000000" not in cu
        assert moi != cu


# ── Nối với hệ thống chung (test đặt trong file NÀY, không thêm vào file test của người khác) ─────────

class TestNoiVoiHeThongChung:
    def test_hai_ban_classify_backend_frontend_ra_cung_ket_qua(self):
        """Frontend giữ bản sao `_classify_upload_filename` (không import backend) — canh không cho hai bản lệch nhau."""
        from frontend.pages.cham_459901_000000000 import _classify_upload_filename as fe
        mau = [
            "GL02_20260731_1000.zip", "gl02_thang8.ZIP", "459_TON_T7.xlsx", "459_ton_thang8.xlsx", "459-mã 0.xlsx",
            "459_mã 0.xlsx", "459 ma 0.xlsx", "459_MÃ_0_T7.xlsx", "459_TON.xls", "459-mã 0.zip", "459-mã 05.xlsx",
            "GL02_20260731_1000.xlsx", "1000_gl02_20260701_459901.xlsx", "1000_gl02_2026080120260831.xlsx",
            "~$459_TON.xlsx", "~$459-mã 0.xlsx", "~$1000_gl02_2026080120260831_1.xlsx", "459_TON.xlsx",
            "GL02 459901 Tổng hợp T8.xlsx", "459901_tong_hop.xlsx", "459901_Boston.xlsx", "459901 - Anton.xlsx",
            "readme.txt", "bao_cao_khac.zip", "khong_nhan_dien_duoc.xlsx",
        ]
        lech = [(t, svc.classify_upload_filename(t), fe(t)) for t in mau if svc.classify_upload_filename(t) != fe(t)]
        assert not lech, f"2 bản classify_upload_filename() lệch nhau (nhãn hiển thị sẽ sai với phân loại thật):\n  {lech}"

    def test_lich_don_23h_that_su_don_thu_muc_cua_module_nay(self, tmp_path, monkeypatch):
        """Gọi `run_cleanup()` thật (không mock): thiếu dòng của module này trong `temp_cleanup_service` thì thư mục
        kết quả không bao giờ bị dọn theo lịch — và TypeError bị nuốt chỉ còn log ERROR."""
        from backend.services import temp_cleanup_service
        monkeypatch.setattr(svc, "TEMP_DIR", tmp_path)
        cu = tmp_path / "ket-qua-cu"
        cu.mkdir()
        mtime = time.time() - 999_999
        os.utime(cu, (mtime, mtime))
        temp_cleanup_service.run_cleanup(cutoff=time.time())
        assert not cu.exists()

    def test_chay_o_tien_trinh_rieng_bao_loi_dung_kieu(self, tien_trinh_that, tmp_path, caplog):
        """Chạy qua tiến trình con THẬT (conftest mặc định chạy trong luồng nên các test khác không đi qua đây)."""
        hong = tmp_path / "GL02_hong.zip"
        hong.write_bytes(b"khong phai zip")
        token = svc.init_progress()
        with caplog.at_level(logging.INFO, logger="backend.core.tien_trinh_doi_chieu"):
            svc.run_process([("GL02_hong.zip", hong)], token)
        p = svc.get_progress(token)
        assert p["done"] and p["error"] and not p["cancelled"]
        assert p["msg"] == p["error"]              # InputError mang về đúng kiểu → thông báo thẳng cho người dùng
        assert p["pct"] >= 5                       # tiến độ từ tiến trình con về tới _progress của cha
        assert any("tiến trình riêng" in r.getMessage() for r in caplog.records)

    def test_nut_dung_toi_duoc_tien_trinh_con(self, tien_trinh_that, tmp_path):
        # Event của cha → Event liên tiến trình → `_set_prog` trong con ném `_Cancelled` → về cha đúng kiểu → "Đã dừng".
        # Sai một khâu là lượt chạy tới hết mà không dừng.
        hong = tmp_path / "GL02.zip"
        hong.write_bytes(b"khong phai zip")
        token = svc.init_progress()
        assert svc.cancel_progress(token)
        svc.run_process([("GL02.zip", hong)], token)
        p = svc.get_progress(token)
        assert p["done"] and p["cancelled"] and not p["error"], p
