# -*- coding: utf-8 -*-
"""Báo cáo dữ liệu thanh toán SWIFT (mẫu D00054) — hai chỗ hỏng lặng lẽ.

1. Dòng TỔNG do công cụ export tự thêm (cột CTHED để trống) phải bị loại TƯỜNG
   MINH. Trước đây nó chỉ rơi ra tình cờ vì tra tên quốc gia không thấy; công cụ
   export mà điền tên vào đó là số liệu nhân đôi, không lỗi, không cảnh báo.
   File thật cho thấy cấu trúc đổi theo kỳ: OUT_202606 không có dòng tổng,
   OUT_202608 có.

2. Quốc gia / vùng lãnh thổ không có dòng trong mẫu (Bermuda, Cayman, và cả
   South Sudan) phải được TRẢ RA cho người dùng, không im lặng bỏ. Đo trên dữ
   liệu thật: kỳ 202606 hụt 59 điện, kỳ 202608 hụt 43 điện mà màn hình không
   báo gì.

Bất biến then chốt: nguồn = báo cáo + bỏ qua. Không được phép rơi vãi chỗ nào.
"""
import io

import openpyxl
import pytest

from backend.services.th_report_service import (
    fill_template, parse_incoming, parse_outgoing,
    _COL_IN_COUNT, _COL_CN_COUNT, _COL_DN_COUNT, _COL_FI_COUNT, _TOTAL_ROW,
)


# ── Dựng file nguồn giả lập ───────────────────────────────────────────────────
def _xlsx(header: list, rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result"
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _file_den(rows) -> bytes:
    """rows: (CTHED, STTLM_AMT, TOTAL)"""
    return _xlsx(["CTHED", "CTRYCD", "STTLM_CUR", "STTLM_AMT", "TOTAL"],
                 [[c, "XX", "USD", a, t] for c, a, t in rows])


def _file_di(rows) -> bytes:
    """rows: (CTHED, CUST_TYPE, TOTAL_AMT, TOTAL)"""
    return _xlsx(["CTHED", "CTRYCD", "CUST_TYPE", "STTLM_CUR", "TOTAL_AMT", "TOTAL"],
                 [[c, "XX", ct, "USD", a, t] for c, ct, a, t in rows])


def _tong_bao_cao(xls: bytes) -> tuple[int, int]:
    ws = openpyxl.load_workbook(io.BytesIO(xls)).active
    den = ws.cell(_TOTAL_ROW, _COL_IN_COUNT).value or 0
    di = sum(ws.cell(_TOTAL_ROW, c).value or 0
             for c in (_COL_CN_COUNT, _COL_DN_COUNT, _COL_FI_COUNT))
    return int(den), int(di)


# ── Dòng tổng ─────────────────────────────────────────────────────────────────
def test_dong_tong_khong_duoc_cong_vao_so_lieu():
    # Dòng cuối là dòng tổng của công cụ export: CTHED trống, TOTAL = tổng phía trên
    den = parse_incoming(_file_den([
        ("JAPAN", "1000", "3"),
        ("FRANCE", "2000", "4"),
        (None, "3000", "7"),
    ]))
    assert den["JAPAN"]["count"] == 3
    assert den["FRANCE"]["count"] == 4
    assert sum(v["count"] for v in den.values()) == 7, "dòng tổng bị cộng vào -> nhân đôi"


def test_dong_tong_khong_lot_vao_canh_bao():
    """Nó không có tên quốc gia, hiện lên bảng cảnh báo là doạ người dùng."""
    den = parse_incoming(_file_den([("JAPAN", "1000", "3"), (None, "1000", "3")]))
    xls, bo_qua = fill_template(den, {}, "202608")
    assert [x["quoc_gia"] for x in bo_qua] == []


def test_dong_tong_file_di_co_cust_type_van_bi_loai():
    """Chốt chặn thật: nhận diện bằng CTHED trống, KHÔNG dựa vào CUST_TYPE rỗng.

    Hiện dòng tổng file OUT có CUST_TYPE rỗng nên bị bỏ qua nhờ nhánh khác —
    công cụ export điền 'DN' vào đó là số liệu đi gấp đôi."""
    di = parse_outgoing(_file_di([
        ("JAPAN", "DN", "1000", "5"),
        (None, "DN", "1000", "5"),
    ]))
    assert sum(v["dn_count"] for v in di.values()) == 5


# ── Quốc gia không có dòng trong mẫu ──────────────────────────────────────────
def test_bao_cao_ra_danh_sach_quoc_gia_bi_bo():
    den = parse_incoming(_file_den([
        ("JAPAN", "1000", "10"),
        ("BERMUDA", "4040", "1"),          # vùng lãnh thổ, mẫu không có dòng
        ("SOUTH SUDAN", "500", "2"),       # quốc gia thật, mẫu D00054 (~2010) thiếu
    ]))
    di = parse_outgoing(_file_di([
        ("JAPAN", "CN", "700", "3"),
        ("CAYMAN ISLANDS", "DN", "1640", "1"),
    ]))
    xls, bo_qua = fill_template(den, di, "202608")

    ten = {x["quoc_gia"]: x for x in bo_qua}
    assert set(ten) == {"BERMUDA", "SOUTH SUDAN", "CAYMAN ISLANDS"}
    assert ten["BERMUDA"]["den"] == 1 and ten["BERMUDA"]["gt_den"] == 4.04
    assert ten["SOUTH SUDAN"]["den"] == 2
    assert ten["CAYMAN ISLANDS"]["di"] == 1 and ten["CAYMAN ISLANDS"]["gt_di"] == 1.64


def test_bao_toan_nguon_bang_bao_cao_cong_bo_qua():
    """Bất biến quan trọng nhất: không điện nào biến mất khỏi cả hai đường."""
    den_rows = [("JAPAN", "1000", "10"), ("BERMUDA", "10", "1"),
                ("GUAM", "10", "4"), (None, "1020", "15")]
    di_rows = [("JAPAN", "CN", "100", "3"), ("FRANCE", "DN", "100", "6"),
               ("REUNION", "TCTD", "100", "2")]
    den = parse_incoming(_file_den(den_rows))
    di = parse_outgoing(_file_di(di_rows))
    xls, bo_qua = fill_template(den, di, "202608")

    rep_den, rep_di = _tong_bao_cao(xls)
    assert rep_den + sum(x["den"] for x in bo_qua) == 15
    assert rep_di + sum(x["di"] for x in bo_qua) == 11


def test_viet_nam_khong_tinh_va_khong_bao_thieu():
    """Việt Nam cố ý để 0 (xem _SKIP_COUNTRIES) — đó không phải dòng bị bỏ sót."""
    den = parse_incoming(_file_den([("VIET NAM", "9999", "1056"), ("JAPAN", "1000", "3")]))
    xls, bo_qua = fill_template(den, {}, "202608")
    assert bo_qua == []
    assert _tong_bao_cao(xls)[0] == 3


def test_khong_thieu_gi_thi_danh_sach_rong():
    den = parse_incoming(_file_den([("JAPAN", "1000", "3")]))
    di = parse_outgoing(_file_di([("FRANCE", "DN", "500", "2")]))
    xls, bo_qua = fill_template(den, di, "202608")
    assert bo_qua == []
    assert _tong_bao_cao(xls) == (3, 2)


# ── Đường truyền cảnh báo về frontend ─────────────────────────────────────────
def test_canh_bao_di_qua_duoc_header_http():
    """Cảnh báo đi kèm thân file Excel nên phải nhét vào header — header HTTP chỉ
    chở latin-1, tên quốc gia phải sống sót qua quote/unquote."""
    import json
    from urllib.parse import quote, unquote

    den = parse_incoming(_file_den([("VIRGIN ISLANDS, BRITISH", "168370", "5")]))
    _, bo_qua = fill_template(den, {}, "202608")

    header = quote(json.dumps(bo_qua, ensure_ascii=True))
    header.encode("latin-1")                      # ném ngay nếu lọt ký tự lạ
    assert json.loads(unquote(header)) == bo_qua
