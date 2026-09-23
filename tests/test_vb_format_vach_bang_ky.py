"""Rà "VB goc" 23/09/2026: đường kẻ có sẵn, bảng Kính gửi, số tự động, khối chữ ký.

Mỗi test dựng văn bản tối giản tái hiện đúng một lỗi đo được trên file thật
(`Dữ liệu test/Chuẩn hoá văn bản/20260923`).
"""
import io
import re

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, Twips

from backend.services.vb_format import ap_dung, bien_doi, do_chu, duong_ke, quy_chuan
from backend.services.vb_format.chuan_hoa import chuan_hoa


def _can_phong_times():
    if do_chu.be_rong_pt("Kính trình:", 14) is None:
        pytest.skip("máy không có Times New Roman để đo bề rộng chữ")


def _luu(doc) -> bytes:
    ra = io.BytesIO()
    doc.save(ra)
    return ra.getvalue()


def _mo(du_lieu: bytes):
    return Document(io.BytesIO(du_lieu))


def _tim(doc, mo_dau: str):
    return next(p for p, _ in ap_dung.duyet_doan(doc) if p.text.strip().startswith(mo_dau))


# ── Đường kẻ tác giả vẽ sẵn ─────────────────────────────────────────────────
_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
       'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
       'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
       'xmlns:v="urn:schemas-microsoft-com:vml"')


def _run_vach(dai_pt: float):
    """Straight Connector y như Word ghi: DrawingML + bản dự phòng VML, neo lệch trái."""
    emu = int(dai_pt * 12700)
    return parse_xml(
        f'<w:r {_NS}><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>862812</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>36830</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{emu}" cy="0"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="1" name="Straight Connector 1"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu}" cy="0"/>'
        '</a:xfrm><a:prstGeom prst="line"><a:avLst/></a:prstGeom></wps:spPr><wps:bodyPr/></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:line id="l1" style="position:absolute;z-index:1;'
        'mso-position-horizontal:absolute;mso-position-horizontal-relative:text" '
        f'from="67.95pt,2.9pt" to="{67.95 + dai_pt:.2f}pt,2.9pt" strokeweight=".5pt"/>'
        '</w:pict></mc:Fallback></mc:AlternateContent></w:r>')


def _van_ban_co_vach_ngan() -> bytes:
    """Tiêu ngữ có sẵn vạch 60 pt — vẽ cho cỡ chữ cũ, nay ngắn hơn dòng chữ."""
    doc = Document()
    doc.add_paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    doc.add_paragraph("Độc lập - Tự do - Hạnh phúc")
    doc.add_paragraph()._p.append(_run_vach(60.0))
    doc.add_paragraph("Hà Nội, ngày 01 tháng 01 năm 2026")
    doc.add_paragraph("Điều 1. Nội dung.")
    return _luu(doc)


def test_vach_co_san_duoc_chinh_bang_dong_chu_va_canh_giua():
    """VB goc: vạch dưới Tiêu ngữ còn 73% dòng chữ sau khi chữ lên cỡ 13 — Điều 7.2 đòi bằng."""
    _can_phong_times()
    du_lieu, bc = chuan_hoa(_van_ban_co_vach_ngan())
    doc = _mo(du_lieu)
    xml = doc.element.body.xml

    mong = do_chu.be_rong_pt("Độc lập - Tự do - Hạnh phúc", 13, True)
    emu = [int(x) for x in re.findall(r'<wp:extent cx="(\d+)"', xml)]
    assert emu and abs(emu[0] / 12700 - mong) < 0.5, (emu, mong)
    # a:ext trong a:xfrm đổi theo — lệch nhau thì Word vẽ theo một trong hai, không đoán được
    assert [int(x) for x in re.findall(r'<a:ext cx="(\d+)"', xml)] == emu
    assert "<wp:align>center</wp:align>" in xml and "posOffset>862812" not in xml
    # Bản dự phòng VML cũng đúng độ dài, canh giữa
    m = re.search(r'from="0,2.9pt" to="([\d.]+)pt,2.9pt"', xml)
    assert m and abs(float(m.group(1)) - mong) < 0.1
    assert "mso-position-horizontal:center" in xml
    # Chỉ một vạch — không vẽ thêm vạch thứ hai
    assert sum(1 for p, _ in ap_dung.duyet_doan(doc) if duong_ke._co_hinh_duong_ke(p._p.xml)) == 1
    assert any("chỉnh độ dài đường kẻ" in v for d in bc["doan"] for v in d["viec"])


def test_vach_co_san_chay_lan_hai_khong_sua_nua():
    _can_phong_times()
    lan1, _ = chuan_hoa(_van_ban_co_vach_ngan())
    lan2, bc = chuan_hoa(lan1)
    assert not any("đường kẻ" in v for d in bc["doan"] for v in d["viec"])
    assert _mo(lan2).element.body.xml.count("<wp:extent") == 1


def test_vach_moi_duoi_ten_don_vi_dai_mot_nua_dong_chu():
    """Điều 8.2: 1/3–1/2. Lấy 1/2 — đúng mức đo trên Phụ lục V (0,50–0,57).
    Người dùng đề nghị 2/3: vượt dải, không lấy."""
    assert duong_ke.TY_LE["ten_dv_ban_hanh"] == 0.5
    assert 1 / 3 <= duong_ke.TY_LE["ten_dv_ban_hanh"] <= 1 / 2
    assert duong_ke.TY_LE["tieu_ngu"] == 1.0


# ── Bảng Kính gửi / Kính trình ──────────────────────────────────────────────
def _van_ban_bang_kinh_trinh(ben_phai: list[str], contextual: bool = False) -> bytes:
    doc = Document()
    doc.add_paragraph("TỜ TRÌNH")
    doc.add_paragraph("V/v thử nghiệm")
    bang = doc.add_table(rows=1, cols=2)
    bang.style = doc.styles["Table Grid"]
    trai, phai = bang.rows[0].cells
    trai.width, phai.width = Twips(1416), Twips(4586)
    trai.paragraphs[0].add_run("Kính trình:")
    phai.paragraphs[0].add_run(ben_phai[0])
    for t in ben_phai[1:]:
        phai.add_paragraph(t)
    if contextual:
        for o in (trai, phai):
            for p in o.paragraphs:
                p._p.get_or_add_pPr().append(parse_xml(f'<w:contextualSpacing {nsdecls("w")}/>'))
    doc.add_paragraph("Căn cứ Quyết định số 01/QĐ ngày 01/01/2026;")
    return _luu(doc)


def _bang_kinh(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text.startswith("Kính"))


def test_bang_kinh_trinh_mot_noi_vua_mot_dong_canh_giua():
    """VB goc: ô 70,8 pt → "Kính / trình:" gãy hai dòng ở cỡ 14, tên người nhận gãy theo."""
    _can_phong_times()
    ten = "Phó Tổng Giám đốc Vương Hồng Lĩnh."
    du_lieu, bc = chuan_hoa(_van_ban_bang_kinh_trinh([ten]))
    tbl = _bang_kinh(_mo(du_lieu))._tbl
    tblPr = tbl.tblPr

    luoi = [int(c.get(qn("w:w"))) for c in tbl.tblGrid.findall(qn("w:gridCol"))]
    assert luoi[0] == round(do_chu.be_rong_pt("Kính trình: ", 14) * 20)
    assert luoi[1] >= do_chu.be_rong_pt(ten.replace(" ", " "), 14) * 20
    assert tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"
    assert tblPr.find(qn("w:tblCellMar")).find(qn("w:left")).get(qn("w:w")) == "0"
    assert all(v.get(qn("w:val")) == "none" for v in tblPr.find(qn("w:tblBorders")))
    assert any("Kính gửi / Kính trình" in s for s in bc["sua_chung"])

    lan2, bc2 = chuan_hoa(du_lieu)
    assert not any("Kính gửi / Kính trình" in s for s in bc2["sua_chung"]), "chạy lại không được báo sửa"


def test_bang_kinh_gui_nhieu_noi_cot_trai_du_dau_hai_cham():
    """Bản đầu cắt dấu ':' khỏi cột trái cho gạch nằm đúng dưới nó — phản biện dựng PDF
    bằng Word: "Kính / trình:" gãy hai dòng. Cột trái phải chứa trọn "Kính trình:"."""
    _can_phong_times()
    du_lieu, _ = chuan_hoa(_van_ban_bang_kinh_trinh(["- Ban Pháp chế;", "- Ban Tài chính Kế toán."]))
    tbl = _bang_kinh(_mo(du_lieu))._tbl
    trai = int(tbl.tblGrid.findall(qn("w:gridCol"))[0].get(qn("w:w")))
    assert trai >= do_chu.be_rong_pt("Kính trình:", 14) * 20
    # nhưng không kèm dấu cách như ca một nơi: gạch đầu dòng bám sát dấu hai chấm
    assert trai < do_chu.be_rong_pt("Kính trình: ", 14) * 20


def test_bang_kinh_gui_co_o_gop_thi_de_nguyen():
    doc = _mo(_van_ban_bang_kinh_trinh(["Giám đốc."]))
    tbl = _bang_kinh(doc)._tbl
    tbl.tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="500"/>'))
    tcPr = tbl.findall(qn("w:tr"))[0].findall(qn("w:tc"))[1].get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:gridSpan {nsdecls("w")} w:val="2"/>'))
    truoc = [c.get(qn("w:w")) for c in tbl.tblGrid.findall(qn("w:gridCol"))]
    du_lieu, _ = chuan_hoa(_luu(doc))
    sau = _bang_kinh(_mo(du_lieu))._tbl
    assert [c.get(qn("w:w")) for c in sau.tblGrid.findall(qn("w:gridCol"))] == truoc


def test_bang_kinh_trinh_ten_qua_dai_khong_vuot_vung_chu():
    _can_phong_times()
    du_lieu, _ = chuan_hoa(_van_ban_bang_kinh_trinh(["Phó Tổng Giám đốc " + "Nguyễn Văn A " * 12]))
    doc = _mo(du_lieu)
    s = doc.sections[0]
    vung_chu = (s.page_width - s.left_margin - s.right_margin) / 635
    w = int(_bang_kinh(doc)._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")))
    assert w <= vung_chu + 1


def test_kinh_trinh_tat_khong_cach_doan_cung_style():
    """VB goc: ô Kính trình có after=6 pt mà vẫn dính "I." — `contextualSpacing` nuốt mất
    (dựng PDF bằng Word: gỡ thẻ là hiện đúng 6 pt)."""
    du_lieu, _ = chuan_hoa(_van_ban_bang_kinh_trinh(["Giám đốc."], contextual=True))
    p = _tim(_mo(du_lieu), "Kính trình")
    cs = p._p.pPr.find(qn("w:contextualSpacing"))
    assert cs is not None and cs.get(qn("w:val")) == "0"
    assert ap_dung._hieu_luc_doan(p, "space_after").pt == 6


def test_loi_van_khong_bi_tat_contextual_spacing():
    """Chỉ thành phần KHAI RIÊNG khoảng cách > 0 mới bị đụng — lời văn giữ nguyên."""
    doc = Document()
    doc.add_paragraph("Điều 1. Nội dung.")
    p = doc.add_paragraph("Lời văn thường của người soạn.")
    p._p.get_or_add_pPr().append(parse_xml(f'<w:contextualSpacing {nsdecls("w")}/>'))
    du_lieu, _ = chuan_hoa(_luu(doc))
    cs = _tim(_mo(du_lieu), "Lời văn")._p.pPr.find(qn("w:contextualSpacing"))
    assert cs is not None and cs.get(qn("w:val")) is None


# ── Một dấu cách sau số ──────────────────────────────────────────────────────
def _van_ban_so_tu_dong(them_bang: bool = False) -> bytes:
    doc = Document()
    doc.add_paragraph("Điều 1. Trách nhiệm thi hành")
    doc.add_paragraph("Căn cứ trình", style="List Number")
    if them_bang:
        bang = doc.add_table(rows=1, cols=2)
        bang.rows[0].cells[0].paragraphs[0].style = doc.styles["List Number"]
        bang.rows[0].cells[0].paragraphs[0].add_run("Số liệu")
    return _luu(doc)


def _lvl_cua(doc, mo_dau: str):
    return ap_dung._tim_lvl(_tim(doc, mo_dau))[1]


def test_so_tu_dong_cach_chu_mot_dau_cach():
    """VB goc: "I." "II." "IV." tự động cách chữ ~0,9 cm, "III." gõ tay cách một dấu cách."""
    du_lieu, bc = chuan_hoa(_van_ban_so_tu_dong())
    doc = _mo(du_lieu)
    lvl = _lvl_cua(doc, "Căn cứ trình")
    assert lvl.find(qn("w:suff")).get(qn("w:val")) == "space"
    # w:suff phải đứng ngay trước w:lvlText — sai chỗ là Word bỏ qua
    ten = [c.tag.split("}")[1] for c in lvl]
    assert ten.index("suff") + 1 == ten.index("lvlText")
    assert len(_tim(doc, "Căn cứ trình").paragraph_format.tab_stops) == 0
    assert any("một dấu cách" in s for s in bc["sua_chung"])


def test_so_tu_dong_dung_chung_voi_bang_so_lieu_thi_giu_tab():
    """`suff` nằm trong định nghĩa danh sách dùng chung — cấp có mặt trong bảng số liệu
    thì không đổi, kẻo lệch cột người soạn canh bằng thụt treo."""
    du_lieu, _ = chuan_hoa(_van_ban_so_tu_dong(them_bang=True))
    assert _lvl_cua(_mo(du_lieu), "Căn cứ trình").find(qn("w:suff")) is None


def test_la_ma_go_tay_trong_loi_van_ve_mot_dau_cach():
    cfg = quy_chuan.mac_dinh()["danh_so"]
    assert bien_doi.chuan_danh_so("III.\tThẩm quyền phê duyệt", "noi_dung", cfg) == [(0, 5, "III. ")]
    assert bien_doi.chuan_danh_so("IV.   Đề xuất", "noi_dung", cfg) == [(0, 6, "IV. ")]
    assert bien_doi.chuan_danh_so("III. Thẩm quyền", "noi_dung", cfg) == []
    # "C. " "D. " đầu câu lời văn: nhiều khả năng là tên viết tắt, không đụng
    assert bien_doi.chuan_danh_so("D.  Nguyễn Văn A đề xuất", "noi_dung", cfg) == []


# ── Khối chữ ký ──────────────────────────────────────────────────────────────
def _van_ban_khoi_ky(so_dong_trong: int = 5, phe_duyet: bool = True) -> bytes:
    doc = Document()
    # Mốc "Số:" chặn luật khối tên đơn vị — thiếu nó thì mọi dòng in hoa trong 12
    # đoạn đầu (kể cả "GIÁM ĐỐC") bị nhận là tên đơn vị.
    doc.add_paragraph("Số: 01/TTr-TTTT")
    doc.add_paragraph("Điều 1. Nội dung.")
    bang = doc.add_table(rows=1, cols=2)
    trai, phai = bang.rows[0].cells
    trai.paragraphs[0].add_run("Nơi nhận:")
    trai.add_paragraph("- Lưu: VT.")
    phai.paragraphs[0].add_run("GIÁM ĐỐC")
    for _ in range(so_dong_trong):
        p = phai.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="16"/><w:lang w:val="vi-VN"/></w:rPr>'))
    phai.add_paragraph("Nguyễn Quốc Hùng")
    if phe_duyet:
        p = doc.add_paragraph("PHÊ DUYỆT CỦA PHÓ TỔNG GIÁM ĐỐC")
        p.paragraph_format.space_before = Pt(12)
    return _luu(doc)


def test_dong_trong_cho_ky_lay_co_chu_khoi_ky():
    """VB goc: 5 dòng trống cỡ 8 dưới "GIÁM ĐỐC" (đã lên cỡ 14) → chỗ ký chỉ ~1,6 cm."""
    du_lieu, bc = chuan_hoa(_van_ban_khoi_ky())
    doc = _mo(du_lieu)
    o = next(t for t in doc.tables).rows[0].cells[1]
    trong = [p for p in o.paragraphs if not p.text.strip()]
    assert len(trong) == 5, "giữ nguyên số dòng tác giả để"
    for p in trong:
        rPr = p._p.pPr.find(qn("w:rPr"))
        assert rPr.find(qn("w:sz")).get(qn("w:val")) == "28"
        # w:sz phải đứng TRƯỚC w:lang theo lược đồ
        ten = [c.tag.split("}")[1] for c in rPr]
        assert ten.index("sz") < ten.index("lang")
        assert ap_dung._hieu_luc_doan(p, "space_after").pt == 0
        assert ap_dung._hieu_luc_doan(p, "line_spacing") == 1.0
    assert any("chừa chữ ký" in s for s in bc["sua_chung"])


def test_khoi_phe_duyet_cach_ho_ten_mot_dong():
    """Mẫu 06: "PHÊ DUYỆT CỦA …" cách "Họ và tên" một dòng. VB goc: before 12 pt bị bỏ về 0
    (luật `bo_khoang_truoc_doan`) → dính sát tên người ký."""
    du_lieu, _ = chuan_hoa(_van_ban_khoi_ky())
    p = _tim(_mo(du_lieu), "PHÊ DUYỆT")
    assert ap_dung._hieu_luc_doan(p, "space_before").pt == pytest.approx(ap_dung.mot_dong_pt(14), abs=0.05)

    lan2, bc = chuan_hoa(du_lieu)
    assert not any("cách khối trên" in v for d in bc["doan"] for v in d["viec"])


def test_chuc_danh_khong_sau_ho_ten_van_bo_khoang_truoc():
    """Chỉ khối ký THỨ HAI (ngay dưới họ tên) mới được giữ khoảng trước."""
    ma = ["noi_dung", "quyen_han_chuc_vu", "trong", "ho_ten_nguoi_ky", "quyen_han_chuc_vu"]
    assert not ap_dung.la_khoi_ky_moi(ma, 1)
    assert ap_dung.la_khoi_ky_moi(ma, 4)


def test_hai_chu_ky_song_song_khong_bi_day_xuong():
    """Phản biện: ô trái "KT. GIÁM ĐỐC … Nguyễn Văn A", ô phải "TRƯỞNG PHÒNG … Trần Văn B"
    cùng một hàng — ô phải từng nhận before 16,1 pt, chữ ký bên phải tụt một dòng."""
    doc = Document()
    doc.add_paragraph("Số: 01/BB-TTTT")
    doc.add_paragraph("Điều 1. Nội dung.")
    trai, phai = doc.add_table(rows=1, cols=2).rows[0].cells
    trai.paragraphs[0].add_run("KT. GIÁM ĐỐC")
    trai.add_paragraph("")
    trai.add_paragraph("Nguyễn Văn An")
    phai.paragraphs[0].add_run("TRƯỞNG PHÒNG")
    phai.add_paragraph("")
    phai.add_paragraph("Trần Văn Bình")
    du_lieu, _ = chuan_hoa(_luu(doc))
    p = _tim(_mo(du_lieu), "TRƯỞNG")          # "TRƯỞNG PHÒNG" đã bị ghép dấu cách không ngắt
    tr = ap_dung._hieu_luc_doan(p, "space_before")
    assert tr is None or tr.pt == 0


def test_khoi_phe_duyet_da_co_dong_trong_thi_khong_cong_them():
    """Tác giả đã để một dòng trống giữa họ tên và "PHÊ DUYỆT…" — cộng thêm là cách hai dòng."""
    ma = ["quyen_han_chuc_vu", "ho_ten_nguoi_ky", "trong", "quyen_han_chuc_vu"]
    assert not ap_dung.la_khoi_ky_moi(ma, 3)
    assert ap_dung.la_khoi_ky_moi(["quyen_han_chuc_vu", "ho_ten_nguoi_ky", "quyen_han_chuc_vu"], 2)


def test_suff_nothing_cua_tac_gia_giu_nguyen():
    """suff="nothing" = tác giả tự gõ khoảng cách trong lvlText ("Điều %1. ") — đổi sang
    "space" là thành hai dấu cách."""
    doc = _mo(_van_ban_so_tu_dong())
    lvl = _lvl_cua(doc, "Căn cứ trình")
    lvl.find(qn("w:lvlText")).addprevious(parse_xml(f'<w:suff {nsdecls("w")} w:val="nothing"/>'))
    du_lieu, _ = chuan_hoa(_luu(doc))
    assert _lvl_cua(_mo(du_lieu), "Căn cứ trình").find(qn("w:suff")).get(qn("w:val")) == "nothing"


def test_hai_vach_cung_mot_doan_thi_khong_dung():
    """Phản biện: hai vạch neo chung một đoạn bằng toạ độ tuyệt đối — ép cả hai cùng độ dài,
    cùng canh giữa một cột là chồng lên nhau. Không chắc là MỘT vạch thì để yên."""
    _can_phong_times()
    doc = _mo(_van_ban_co_vach_ngan())
    p_vach = next(p for p, _ in ap_dung.duyet_doan(doc) if duong_ke._co_hinh_duong_ke(p._p.xml))
    p_vach._p.append(_run_vach(40.0))
    du_lieu, bc = chuan_hoa(_luu(doc))
    emu = sorted(int(x) for x in re.findall(r'<wp:extent cx="(\d+)"', _mo(du_lieu).element.body.xml))
    assert emu == [40 * 12700, 60 * 12700]
    assert not any("chỉnh độ dài đường kẻ" in v for d in bc["doan"] for v in d["viec"])


def test_dong_trong_cho_ky_co_track_changes_rpr_dung_truoc_pprchange():
    doc = _mo(_van_ban_khoi_ky(so_dong_trong=1, phe_duyet=False))
    o = next(t for t in doc.tables).rows[0].cells[1]
    p = next(p for p in o.paragraphs if not p.text.strip())
    pPr = p._p.get_or_add_pPr()
    pPr.remove(pPr.find(qn("w:rPr")))
    pPr.append(parse_xml(f'<w:pPrChange {nsdecls("w")} w:id="1" w:author="A" '
                         'w:date="2026-01-01T00:00:00Z"><w:pPr/></w:pPrChange>'))
    du_lieu, _ = chuan_hoa(_luu(doc))
    o = next(t for t in _mo(du_lieu).tables).rows[0].cells[1]
    p = next(p for p in o.paragraphs if not p.text.strip())
    ten = [c.tag.split("}")[1] for c in p._p.pPr]
    assert ten.index("rPr") < ten.index("pPrChange")
