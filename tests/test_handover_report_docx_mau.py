"""File Word báo cáo bàn giao phải giữ đúng 4 điểm người dùng chốt 09/09/2026.

Bốn điểm này đều là thứ **không nhìn ra khi đọc code gọi hàm**: mất cái nào thì
file vẫn xuất ra bình thường, chỉ khác lúc in. Nên phải canh bằng test.
"""
import io

from docx import Document

from backend.services.handover_report_docx import build_report_docx

DATA = {
    "overall": {"total": 10, "on_time": 8, "late": 2, "rate": 80.0},
    "by_dept": [{"dept_name": "Phòng Kế toán", "total": 10, "on_time": 8, "late": 2, "rate": 80.0}],
    "late_entries": [
        {"dept_name": "Phòng Kế toán", "staff_id": 1, "staff_name": "Trịnh Nam Sơn",
         "transaction_date": "2026-08-03", "submitted_date": "2026-08-06",
         "days_late": 2, "sheet_count": 79},
        {"dept_name": "Phòng Kế toán", "staff_id": 1, "staff_name": "Trịnh Nam Sơn",
         "transaction_date": "2026-08-05", "submitted_date": "2026-08-07",
         "days_late": 1, "sheet_count": 58},
    ],
}


def _doc():
    return Document(io.BytesIO(build_report_docx(DATA, 2026, 8)))


def test_moi_bang_lap_lai_dong_tieu_de_khi_sang_trang():
    for tbl in _doc().tables[:-1]:          # bảng cuối là ô ký, không có tiêu đề
        assert "tblHeader" in tbl.rows[0]._tr.xml


def test_so_trang_o_dau_trang_va_bat_dau_tu_trang_2():
    sec = _doc().sections[0]
    assert sec.different_first_page_header_footer is True
    assert " PAGE " in sec.header.paragraphs[0]._p.xml
    # Trang đầu phải có đầu trang RIÊNG và rỗng, không thì Word kéo số trang về
    assert sec.first_page_header.is_linked_to_previous is False
    assert "".join(p.text for p in sec.first_page_header.paragraphs) == ""
    # Chân trang phải sạch — không để sót số trang cũ ở dưới
    assert " PAGE " not in sec.footer.paragraphs[0]._p.xml


def test_o_ky_chi_co_nhan_khong_co_ten():
    sig = _doc().tables[-1]
    assert [c.paragraphs[0].text for c in sig.rows[0].cells] == ["LẬP BẢNG", "KIỂM SOÁT"]
    assert len(sig.rows) == 1
    # Bản mẫu giấy CÓ sẵn 2 tên dưới nhãn; người sau nhìn mẫu rồi in tên lại là
    # chuyện dễ xảy ra. Chỉ kiểm dòng nhãn thì tên lọt qua sạch sẽ.
    assert all(para.text == "" for c in sig.rows[0].cells for para in c.paragraphs[1:])
    # Khối ký không được cắt ngang trang
    assert "cantSplit" in sig.rows[0]._tr.xml


def test_truong_page_tach_lam_ba_run():
    """Word tự ghi field code thành 3 run; gộp lại một run có bản in ra chữ "PAGE"."""
    para = _doc().sections[0].header.paragraphs[0]
    kinds = [
        "begin" if "fldCharType=\"begin\"" in r._r.xml
        else "instr" if "instrText" in r._r.xml
        else "end" if "fldCharType=\"end\"" in r._r.xml
        else "?"
        for r in para.runs
    ]
    assert kinds == ["begin", "instr", "end"]


def test_tieu_de_phong_khong_kem_so_dem_va_cot_ghi_nop_cham():
    doc = _doc()
    texts = [p.text for p in doc.paragraphs]
    assert "1. Phòng Kế toán" in texts
    assert not any("chứng từ quá hạn" in t and t.startswith("1.") for t in texts)
    assert [c.text for c in doc.tables[0].rows[0].cells][4] == "Nộp chậm"
